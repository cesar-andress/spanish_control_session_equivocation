"""Ingest binary-validation N=12 returns; apply locked PASS/FAIL before discussion.

Not Calibration Round 3. Does not alter Round-2 labels. Does not issue main packets.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import date, datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document

from scse.calibration_metrics import _cohens_kappa
from scse.paths import DATA_PRIVATE, PROJECT_ROOT, PROTOCOL_DIR

ROOT = PROJECT_ROOT
VAL = DATA_PRIVATE / "validation" / "binary_v1"
PKT = VAL / "BINARY_VALIDATION_N12_PACKET_BLINDED.csv"
RET_DIR = ROOT / "_internal" / "binary_validation" / "returns"
REPORT_MD = ROOT / "_internal" / "reports" / "BINARY_VALIDATION_AGREEMENT.md"
REPORT_JSON = ROOT / "_internal" / "reports" / "BINARY_VALIDATION_AGREEMENT.json"
GATE_JSON = VAL / "BINARY_VALIDATION_PROSPECTIVE_GATE.json"
INSTRUMENT = PROTOCOL_DIR / "direct_resolution_codebook_v1.0-validation.md"

LABELS = [
    ("direct_resolution_yes", "Sí — resolución directa"),
    ("direct_resolution_no", "No — no resolución directa"),
]
ES = {
    "direct_resolution_yes": "Sí — resolución directa",
    "direct_resolution_no": "No — no resolución directa",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def is_marked(line: str) -> bool:
    s = line.strip()
    if re.match(r"^[xX]\s+", s):
        return True
    return s.startswith(("☑", "☒", "✓", "✔"))


def parse_case_block(text: str, case_no: int) -> dict:
    lines = [ln.rstrip() for ln in text.splitlines()]
    reply_marks: dict[str, bool] = {}
    dudoso_marks: dict[str, bool] = {}
    obs_lines: list[str] = []
    mode = None
    for ln in lines:
        low = ln.strip().lower()
        if low.startswith("tu valoración") or low.startswith("tu valoracion"):
            mode = "reply"
            continue
        if "dudoso" in low:
            mode = "dudoso"
            continue
        if low.startswith("observaciones"):
            mode = "obs"
            continue
        if mode == "reply":
            for canon, needle in LABELS:
                if needle.lower() in low:
                    reply_marks[canon] = is_marked(ln)
        elif mode == "dudoso":
            if re.search(r"^[xX☐☑☒✓✔]\s*sí\b", low) or re.search(
                r"^[xX☐☑☒✓✔]\s*si\b", low
            ):
                dudoso_marks["sí"] = is_marked(ln)
            elif re.search(r"^[xX☐☑☒✓✔]\s*no\b", low):
                dudoso_marks["no"] = is_marked(ln)
        elif mode == "obs":
            if set(ln.strip()) <= {"_", " ", "─", "-", ""}:
                continue
            if ln.strip():
                obs_lines.append(ln.strip())
    selected = [k for k, v in reply_marks.items() if v]
    selected_dudoso = [k for k, v in dudoso_marks.items() if v]
    issues: list[str] = []
    if len(selected) == 0:
        issues.append("missing_direct_resolution")
    elif len(selected) > 1:
        issues.append("multiple_direct_resolution")
    if len(selected_dudoso) == 0:
        issues.append("missing_dudoso")
    elif len(selected_dudoso) > 1:
        issues.append("multiple_dudoso")
    return {
        "case_no": case_no,
        "direct_resolution": selected[0] if len(selected) == 1 else None,
        "direct_resolution_all_marked": selected,
        "dudoso": selected_dudoso[0] if len(selected_dudoso) == 1 else None,
        "observaciones": "\n".join(obs_lines).strip(),
        "issues": issues,
    }


def extract_doc(path: Path) -> list[dict]:
    doc = Document(path)
    case_nos: list[int] = []
    for p in doc.paragraphs:
        m = re.match(r"^CASO\s+(\d+)\s*$", p.text.strip())
        if m:
            case_nos.append(int(m.group(1)))
    if len(doc.tables) != 12 or case_nos != list(range(1, 13)):
        raise RuntimeError(
            f"bad structure {path.name}: cases={case_nos} tables={len(doc.tables)}"
        )
    return [
        parse_case_block(table.cell(0, 0).text, case_no)
        for case_no, table in zip(case_nos, doc.tables)
    ]


def archive(src: Path, dest_name: str) -> dict:
    RET_DIR.mkdir(parents=True, exist_ok=True)
    dest = RET_DIR / dest_name
    if dest.is_file() and sha256(dest) != sha256(src):
        pre = RET_DIR / f"pre_{date.today().isoformat()}_{dest_name}"
        shutil.copy2(dest, pre)
    shutil.copy2(src, dest)
    return {
        "path": str(dest.relative_to(ROOT)),
        "sha256": sha256(dest),
        "source_filename": src.name,
    }


def apply_locked_gate(*, n_agree: int, n: int, kappa: float) -> dict:
    gate_spec = json.loads(GATE_JSON.read_text(encoding="utf-8"))
    k_thr = float(gate_spec["kappa_threshold"])
    raw_min = int(gate_spec["raw_agreement_threshold_n"])
    raw_of = int(gate_spec["raw_agreement_threshold_of"])
    assert n == raw_of == 12
    raw_ok = n_agree >= raw_min
    kappa_ok = kappa >= k_thr
    passed = raw_ok and kappa_ok
    return {
        "gate": "PASS" if passed else "FAIL",
        "passed": passed,
        "kappa": kappa,
        "kappa_threshold": k_thr,
        "kappa_ok": kappa_ok,
        "n_agree": n_agree,
        "n": n,
        "raw_agreement": n_agree / n,
        "raw_agreement_threshold": f"{raw_min}/{raw_of}",
        "raw_ok": raw_ok,
        "rule_locked_before_returns": True,
        "thresholds_not_optimised_after_returns": True,
        "automatic_further_round_on_fail": False,
        "action": (
            "freeze_direct_resolution_v1.0; main_packets_may_be_issued_next"
            if passed
            else "STOP; no_automatic_further_round; no_threshold_change; no_AI_rescue"
        ),
    }


def maybe_freeze_v1(gate: dict, instr_sha: str) -> str | None:
    if not gate["passed"]:
        return None
    src = INSTRUMENT
    dest = PROTOCOL_DIR / "direct_resolution_codebook_v1.0.md"
    text = src.read_text(encoding="utf-8")
    # Editorial freeze header only
    header = (
        "# Guía de criterios: resolución directa de la demanda comunicativa v1.0\n\n"
        "**Estado: CONGELADO tras validación binaria OOP N=12 (PASS)**  \n"
        f"**Fecha de congelación:** {date.today().isoformat()}  \n"
        f"**Instrumento de validación (SHA-256):** `{instr_sha}`  \n"
        f"**Validación:** κ = {gate['kappa']:.3f} (≥ 0,67); "
        f"acuerdo bruto = {gate['n_agree']}/{gate['n']} (≥ 10/12)  \n\n"
        "Derivado de `direct_resolution_codebook_v1.0-validation.md` sin cambiar "
        "reglas sustantivas.\n\n"
        "---\n\n"
    )
    # Drop validation-specific header from source; keep body from first ## 
    body = text
    if "## Linaje" in text:
        body = text[text.index("## Linaje") :]
    elif "## 0." in text:
        body = text[text.index("## 0.") :]
    # Replace validation section 8 with freeze note
    body = re.sub(
        r"## 8\. Instrumento de validación.*",
        "## 8. Congelación\n\n"
        "Instrumento congelado tras PASS de la validación binaria OOP (N=12). "
        "Gobierna la anotación principal XIV. No modificar sin nueva decisión "
        "de protocolo.\n",
        body,
        count=1,
        flags=re.S,
    )
    dest.write_text(header + body, encoding="utf-8")
    return str(dest.relative_to(ROOT))


def main() -> dict:
    ap = argparse.ArgumentParser()
    ap.add_argument("--daniel", type=Path, required=True)
    ap.add_argument("--jose-jaime", type=Path, required=True)
    args = ap.parse_args()

    pkt = pd.read_csv(PKT)
    assert list(pkt["case_no"]) == list(range(1, 13))

    d_meta = archive(args.daniel, "validacion_binaria_daniel_returned.docx")
    j_meta = archive(args.jose_jaime, "validacion_binaria_jose_jaime_returned.docx")
    d_cases = extract_doc(RET_DIR / "validacion_binaria_daniel_returned.docx")
    j_cases = extract_doc(RET_DIR / "validacion_binaria_jose_jaime_returned.docx")

    d_by = {c["case_no"]: c for c in d_cases}
    j_by = {c["case_no"]: c for c in j_cases}

    rows = []
    for _, prow in pkt.iterrows():
        cno = int(prow["case_no"])
        d, j = d_by[cno], j_by[cno]
        complete = d["direct_resolution"] is not None and j["direct_resolution"] is not None
        rows.append(
            {
                "case_no": cno,
                "unit_id": prow["unit_id"],
                "daniel_direct_resolution": d["direct_resolution"],
                "jose_jaime_direct_resolution": j["direct_resolution"],
                "agreement": bool(
                    complete and d["direct_resolution"] == j["direct_resolution"]
                ),
                "pair_complete": complete,
                "daniel_dudoso": d["dudoso"],
                "jose_jaime_dudoso": j["dudoso"],
                "daniel_observaciones": d["observaciones"] or None,
                "jose_jaime_observaciones": j["observaciones"] or None,
                "daniel_issues": ";".join(d["issues"]) or None,
                "jose_jaime_issues": ";".join(j["issues"]) or None,
            }
        )
    codes = pd.DataFrame(rows)
    codes_path = VAL / "BINARY_VALIDATION_CODES.csv"
    codes.to_csv(codes_path, index=False)

    complete = codes[codes["pair_complete"]].copy()
    n = len(complete)
    if n != 12:
        incomplete = codes.loc[~codes["pair_complete"], "case_no"].tolist()
        raise RuntimeError(f"incomplete pairs: {incomplete}")

    a = list(complete["daniel_direct_resolution"])
    b = list(complete["jose_jaime_direct_resolution"])
    n_agree = sum(x == y for x, y in zip(a, b))
    kappa = _cohens_kappa(
        a, b, ["direct_resolution_yes", "direct_resolution_no"]
    )
    disagree = [
        int(r.case_no)
        for r in complete.itertuples()
        if r.daniel_direct_resolution != r.jose_jaime_direct_resolution
    ]
    # confusion rows=Daniel cols=JJ
    labs = ["direct_resolution_yes", "direct_resolution_no"]
    cm = {r: {c: 0 for c in labs} for r in labs}
    for x, y in zip(a, b):
        cm[x][y] += 1

    gate = apply_locked_gate(n_agree=n_agree, n=n, kappa=kappa)
    instr_sha = sha256(INSTRUMENT)
    v1_path = maybe_freeze_v1(gate, instr_sha)

    returns_payload = {
        "phase": "7B_returns",
        "task": "fresh_binary_validation_returns_ingested",
        "not_calibration_round_3": True,
        "recorded_on": date.today().isoformat(),
        "recorded_at": datetime.now(timezone.utc).isoformat(),
        "instrument_path": str(INSTRUMENT.relative_to(ROOT)),
        "instrument_sha256": instr_sha,
        "returned_files": {"daniel": d_meta, "jose_jaime": j_meta},
        "n_cases": 12,
        "n_complete_pairs": n,
        "n_agree": n_agree,
        "n_disagree": n - n_agree,
        "disagree_case_nos": disagree,
        "raw_agreement": n_agree / n,
        "cohen_kappa_binary": kappa,
        "confusion_matrix": cm,
        "locked_gate_decision": gate,
        "codebook_v1_frozen_path": v1_path,
        "main_packets_issued": False,
        "note": (
            "Agreement computed before discussion. Locked PASS/FAIL applied. "
            "Not XIV main outcomes. Not Round 3."
        ),
    }
    (VAL / "BINARY_VALIDATION_RETURNS.json").write_text(
        json.dumps(returns_payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    REPORT_JSON.write_text(
        json.dumps(returns_payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    lines = [
        "# Binary validation agreement (fresh OOP N=12)",
        "",
        f"**Recorded:** {date.today().isoformat()}",
        "**Not** Calibration Round 3.",
        f"**Complete pairs:** {n} / 12",
        "",
        "## Metrics (before discussion)",
        "",
        "| Metric | Value |",
        "|--------|------:|",
        f"| Raw agreement | {n_agree}/{n} = {n_agree/n:.3f} |",
        f"| Cohen κ (binary) | {kappa:.3f} |",
        f"| Agree / disagree | {n_agree} / {n - n_agree} |",
        "",
        f"Disagreement cases: {disagree or '—'}",
        "",
        "## Locked PASS/FAIL",
        "",
        f"**Gate: {gate['gate']}**",
        "",
        f"- κ = {kappa:.3f} (threshold {gate['kappa_threshold']}) → "
        f"{'OK' if gate['kappa_ok'] else 'FAIL'}",
        f"- raw = {n_agree}/{n} (threshold {gate['raw_agreement_threshold']}) → "
        f"{'OK' if gate['raw_ok'] else 'FAIL'}",
        f"- Action: `{gate['action']}`",
        "",
        "Thresholds locked before returns; not optimised after seeing results.",
        "",
        "## Confusion matrix (rows=Daniel, cols=Jose Jaime)",
        "",
        "|  | Sí | No |",
        "|---|---:|---:|",
        f"| Sí | {cm['direct_resolution_yes']['direct_resolution_yes']} | "
        f"{cm['direct_resolution_yes']['direct_resolution_no']} |",
        f"| No | {cm['direct_resolution_no']['direct_resolution_yes']} | "
        f"{cm['direct_resolution_no']['direct_resolution_no']} |",
        "",
        f"Codebook v1.0 frozen: {v1_path or 'no (FAIL)'}",
        "",
    ]
    REPORT_MD.write_text("\n".join(lines), encoding="utf-8")

    return {
        "gate": gate,
        "n_agree": n_agree,
        "kappa": kappa,
        "disagree": disagree,
        "v1": v1_path,
        "confusion": cm,
    }


if __name__ == "__main__":
    print(json.dumps(main(), indent=2, ensure_ascii=False))
