"""Ingest Round-2 independent returns; apply locked A/B/C decision rule.

Does not alter Round-2 draw/seed. Does not invent labels.
Agreement before discussion. Thresholds were locked before returns.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from datetime import date, datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document

from scse.calibration_metrics import ES_TO_EN, summarize_round1
from scse.paths import DATA_PRIVATE, PROJECT_ROOT, PROTOCOL_DIR

ROOT = PROJECT_ROOT
PACKET = DATA_PRIVATE / "calibration" / "round2" / "CALIBRATION_ROUND2_PACKET_BLINDED.csv"
OUT_PRIV = DATA_PRIVATE / "calibration" / "round2"
RET_DIR = ROOT / "_internal" / "calibration_round2" / "returns"
REPORT_MD = ROOT / "_internal" / "reports" / "CALIBRATION_ROUND2_AGREEMENT.md"
REPORT_JSON = ROOT / "_internal" / "reports" / "CALIBRATION_ROUND2_AGREEMENT.json"
JOINT_JSON = RET_DIR / "JOINT_REVIEW_RETURN.json"

ALPHA_THRESHOLD = 0.67
BINARY_KAPPA_THRESHOLD = 0.67

LABELS = [
    ("respuesta explícita", "Respuesta explícita"),
    ("respuesta parcial o intermedia", "Respuesta parcial o intermedia"),
    ("ausencia de respuesta", "Ausencia de respuesta"),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def is_marked(line: str) -> bool:
    s = line.strip()
    if re.match(r"^[xX]\s+", s):
        return True
    return s.startswith(("☑", "☒", "✓", "✔"))


def parse_case_block(text: str, case_no: int) -> dict:
    lines = [ln.rstrip() for ln in text.splitlines()]
    reply_marks: dict[str, bool] = {}
    dudoso_marks: dict[str, bool] = {}
    obs_lines: list[str] = []
    mode = None
    for ln in lines:
        low = ln.strip().lower()
        if low.startswith("tu valoración"):
            mode = "reply"
            continue
        if "dudoso" in low:
            mode = "dudoso"
            continue
        if low.startswith("observaciones"):
            mode = "obs"
            continue
        if mode == "reply":
            for canon, needle in LABELS:
                if needle.lower() in low:
                    reply_marks[canon] = is_marked(ln)
        elif mode == "dudoso":
            if re.search(r"^[xX☐☑☒✓✔]\s*sí\b", low) or re.search(
                r"^[xX☐☑☒✓✔]\s*si\b", low
            ):
                dudoso_marks["sí"] = is_marked(ln)
            elif re.search(r"^[xX☐☑☒✓✔]\s*no\b", low):
                dudoso_marks["no"] = is_marked(ln)
        elif mode == "obs":
            if set(ln.strip()) <= {"_", " ", "─", "-", ""}:
                continue
            if ln.strip():
                obs_lines.append(ln.strip())
    selected_reply = [k for k, v in reply_marks.items() if v]
    selected_dudoso = [k for k, v in dudoso_marks.items() if v]
    issues: list[str] = []
    if len(selected_reply) == 0:
        issues.append("missing_reply_status")
    elif len(selected_reply) > 1:
        issues.append("multiple_reply_status")
    if len(selected_dudoso) == 0:
        issues.append("missing_dudoso")
    elif len(selected_dudoso) > 1:
        issues.append("multiple_dudoso")
    return {
        "case_no": case_no,
        "reply_status": selected_reply[0] if len(selected_reply) == 1 else None,
        "reply_status_all_marked": selected_reply,
        "dudoso": selected_dudoso[0] if len(selected_dudoso) == 1 else None,
        "dudoso_all_marked": selected_dudoso,
        "observaciones": "\n".join(obs_lines).strip(),
        "issues": issues,
    }


def extract_doc(path: Path) -> list[dict]:
    doc = Document(path)
    case_nos: list[int] = []
    for p in doc.paragraphs:
        m = re.match(r"^CASO\s+(\d+)\s*$", p.text.strip())
        if m:
            case_nos.append(int(m.group(1)))
    if len(doc.tables) != 20 or case_nos != list(range(1, 21)):
        raise RuntimeError(
            f"bad structure {path.name}: cases={case_nos} tables={len(doc.tables)}"
        )
    return [
        parse_case_block(table.cell(0, 0).text, case_no)
        for case_no, table in zip(case_nos, doc.tables)
    ]


def apply_locked_rule(metrics: dict) -> dict:
    alpha = float(metrics["krippendorff_alpha_nominal"])
    bin_k = float(metrics["cohen_kappa_explicit_vs_rest"])
    if alpha >= ALPHA_THRESHOLD:
        gate = "A"
        action = (
            "retain_three_category_primary; no_round3; write_v0.4; freeze; "
            "main_annotation_may_proceed_after_packets"
        )
    elif bin_k >= BINARY_KAPPA_THRESHOLD:
        gate = "B"
        action = (
            "binary_explicit_primary; three_category_secondary; no_round3; "
            "write_v0.4; freeze_with_binary_primary_note"
        )
    else:
        gate = "C"
        action = "STOP_before_main; no_automatic_round3; reconsider_scheme"
    return {
        "gate": gate,
        "action": action,
        "alpha_threshold": ALPHA_THRESHOLD,
        "binary_kappa_threshold": BINARY_KAPPA_THRESHOLD,
        "krippendorff_alpha_nominal": alpha,
        "cohen_kappa_explicit_vs_rest": bin_k,
        "rule_locked_before_returns": True,
        "thresholds_not_optimised_after_returns": True,
    }


def write_agreement_md(payload: dict) -> None:
    m = payload["agreement_complete_pairs"]
    gate = payload["locked_rule_decision"]
    lines = [
        "# Calibration Round 2 agreement",
        "",
        f"**Recorded:** {payload['recorded_on']}",
        f"**Codebook at coding:** v0.3",
        f"**Complete pairs:** {payload['n_complete_pairs']} / 20",
        "",
        "## Metrics (before discussion)",
        "",
        "| Metric | Value |",
        "|--------|------:|",
        f"| Raw 3-category agreement | {m['raw_agreement']:.3f} |",
        f"| Cohen κ (3-way) | {m['cohen_kappa_3way']:.3f} |",
        f"| Krippendorff α (nominal) | {m['krippendorff_alpha_nominal']:.3f} |",
        f"| Binary explicit agreement | {m['binary_explicit_raw_agreement']:.3f} |",
        f"| Binary explicit Cohen κ | {m['cohen_kappa_explicit_vs_rest']:.3f} |",
        f"| Agree / disagree (complete) | {m['n_agree']} / {m['n_disagree']} |",
        "",
        f"Disagreement cases: {m['disagree_case_nos'] or '—'}",
        "",
        "## Locked A/B/C decision",
        "",
        f"**Gate: {gate['gate']}**",
        "",
        f"- α = {gate['krippendorff_alpha_nominal']:.3f} (threshold {gate['alpha_threshold']})",
        f"- binary κ = {gate['cohen_kappa_explicit_vs_rest']:.3f} "
        f"(threshold {gate['binary_kappa_threshold']})",
        f"- Action: `{gate['action']}`",
        "",
        "Thresholds were locked before returns; not optimised after seeing results.",
        "",
        "## Confusion matrix (rows=Daniel, cols=Jose Jaime)",
        "",
    ]
    mat = m["confusion_matrix_3way"]
    labs = list(ES_TO_EN.values())
    lines.append("|  | " + " | ".join(labs) + " |")
    lines.append("|---|" + "|".join(["---"] * len(labs)) + "|")
    for r in labs:
        lines.append("| " + r + " | " + " | ".join(str(mat[r][c]) for c in labs) + " |")
    lines += ["", "## Note", "", payload["note"], ""]
    REPORT_MD.write_text("\n".join(lines), encoding="utf-8")


def write_codebook_v04(gate: str) -> Path | None:
    if gate == "C":
        return None
    src = PROTOCOL_DIR / "reply_status_codebook_v0.3.md"
    joint = json.loads(JOINT_JSON.read_text(encoding="utf-8"))
    clar = joint["boundary_clarifications"]
    text = src.read_text(encoding="utf-8")
    # Header swap
    text = text.replace(
        "**Estado: VERSIÓN DE CALIBRACIÓN — BORRADOR (no congelado para el estudio principal)**",
        "**Estado: CONGELADO PARA EL ESTUDIO PRINCIPAL (v0.4)**",
    )
    text = text.replace("**Versión:** `0.3.0`", "**Versión:** `0.4.0`")
    text = text.replace(
        "**Fecha:** 2026-10-05",
        f"**Fecha:** {date.today().isoformat()}",
    )
    text = text.replace(
        "Instrumento de trabajo para la **Ronda 1 de calibración**.",
        "Instrumento congelado tras la **Ronda 2 de calibración** "
        f"(regla bloqueada: puerta **{gate}**).",
    )
    # Append clarifications section faithfully
    appendix = f"""

---

## 12. Aclaraciones de frontera (Ronda 1 conjunta → v0.4)

Texto de los profesores, transcrito sin reescritura substantiva
(revisión conjunta de los casos 5, 6, 7 y 10).

### 12.1 Ausencia de respuesta → respuesta parcial

> {clar['absence_to_partial']}

### 12.2 Respuesta parcial → respuesta explícita

> {clar['partial_to_explicit']}

### 12.3 Decisiones conjuntas de los cuatro casos (contexto de calibración)

Estas marcas no son etiquetas de oro para el estudio principal; documentan el
acuerdo tras discusión de desacuerdos de la Ronda 1.

| Caso | Decisión conjunta |
|------|-------------------|
"""
    for c in joint["joint_case_decisions"]:
        appendix += f"| {c['case_no']} | {c['joint_decision']} |\n"
    if gate == "B":
        appendix += """
### 12.4 Endpoint primario analítico (puerta B)

La fiabilidad tricotómica no alcanzó α ≥ 0.67; la κ binaria
explícita / no-explícita sí alcanzó el umbral bloqueado. El endpoint
**primario** del análisis principal es **explícita vs no-explícita**.
La distribución de tres categorías se conserva y se reporta como
secundaria/descriptiva, con su fiabilidad más débil hecha explícita.
"""
    else:
        appendix += """
### 12.4 Esquema primario (puerta A)

Se mantiene el esquema de tres categorías como esquema primario de
anotación y análisis, tras α nominal ≥ 0.67 en la Ronda 2.
"""
    out = PROTOCOL_DIR / "reply_status_codebook_v0.4.md"
    out.write_text(text.rstrip() + appendix + "\n", encoding="utf-8")
    # changes note
    changes = PROTOCOL_DIR / "reply_status_codebook_v0.3_to_v0.4_changes.md"
    changes.write_text(
        f"""# Codebook v0.3 → v0.4

**Date:** {date.today().isoformat()}  
**Gate:** {gate} (locked Round-2 rule applied after returns)

## What changed

- Status: calibration draft → frozen for main study (v0.4.0).
- Added §12 with joint Round-1 boundary clarifications (verbatim professors’ wording).
- Recorded Round-2 gate {gate} implication for primary analytic scheme.
- v0.3 file left unchanged.

## What did not change

- Core three-category definitions (explicit / partial / absence).
- Unit of analysis; Q1 target; registered anchor.
- No XIV reply_status labels created by this freeze.
""",
        encoding="utf-8",
    )
    return out


def ingest(daniel_docx: Path, jose_docx: Path, joint_docx: Path | None = None) -> dict:
    RET_DIR.mkdir(parents=True, exist_ok=True)
    OUT_PRIV.mkdir(parents=True, exist_ok=True)

    arch: dict = {}
    dest_d = RET_DIR / "calibracion_ronda2_daniel_returned.docx"
    dest_j = RET_DIR / "calibracion_ronda2_jose_jaime_returned.docx"
    shutil.copy2(daniel_docx, dest_d)
    shutil.copy2(jose_docx, dest_j)
    arch["daniel"] = {
        "path": str(dest_d.relative_to(ROOT)),
        "sha256": sha256(dest_d),
        "source_filename": daniel_docx.name,
    }
    arch["jose_jaime"] = {
        "path": str(dest_j.relative_to(ROOT)),
        "sha256": sha256(dest_j),
        "source_filename": jose_docx.name,
    }
    if joint_docx and joint_docx.is_file():
        dest_joint = RET_DIR / "01_REVISION_CONJUNTA_CASOS_5_6_7_10_returned.docx"
        # keep prior archive if identical; else write dated copy + replace canonical
        new_sha = sha256(joint_docx)
        if dest_joint.is_file() and sha256(dest_joint) != new_sha:
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
            shutil.copy2(dest_joint, RET_DIR / f"01_REVISION_CONJUNTA_CASOS_5_6_7_10_returned_pre_{stamp}.docx")
        shutil.copy2(joint_docx, dest_joint)
        arch["joint_review"] = {
            "path": str(dest_joint.relative_to(ROOT)),
            "sha256": sha256(dest_joint),
            "source_filename": joint_docx.name,
        }

    extracted = {
        "daniel": extract_doc(daniel_docx),
        "jose_jaime": extract_doc(jose_docx),
    }
    pkt = pd.read_csv(PACKET).sort_values("case_no")
    unit_by_case = dict(zip(pkt.case_no.astype(int), pkt.unit_id.astype(str)))

    rows = []
    for i in range(20):
        d = extracted["daniel"][i]
        j = extracted["jose_jaime"][i]
        rows.append(
            {
                "case_no": i + 1,
                "unit_id": unit_by_case[i + 1],
                "daniel_reply_status": d["reply_status"],
                "jose_jaime_reply_status": j["reply_status"],
                "agreement_reply_status": bool(
                    d["reply_status"]
                    and j["reply_status"]
                    and d["reply_status"] == j["reply_status"]
                ),
                "pair_complete": d["reply_status"] is not None
                and j["reply_status"] is not None,
                "daniel_dudoso": d["dudoso"],
                "jose_jaime_dudoso": j["dudoso"],
                "daniel_observaciones": d["observaciones"],
                "jose_jaime_observaciones": j["observaciones"],
                "daniel_issues": ";".join(d["issues"]),
                "jose_jaime_issues": ";".join(j["issues"]),
                "daniel_all_marked": "|".join(d["reply_status_all_marked"]),
                "jose_jaime_all_marked": "|".join(j["reply_status_all_marked"]),
            }
        )
    df = pd.DataFrame(rows)
    complete = df[df.pair_complete]
    if len(complete) == 0:
        raise RuntimeError("no complete Round-2 pairs")
    metrics = summarize_round1(
        complete["daniel_reply_status"].tolist(),
        complete["jose_jaime_reply_status"].tolist(),
    )
    metrics["n_complete_pairs"] = int(len(complete))
    metrics["n_incomplete_pairs"] = int((~df.pair_complete).sum())
    metrics["n_agree_complete"] = int(complete["agreement_reply_status"].sum())
    metrics["n_disagree_complete"] = int((~complete["agreement_reply_status"]).sum())
    metrics["disagree_case_nos_complete"] = (
        complete.loc[~complete["agreement_reply_status"], "case_no"].astype(int).tolist()
    )
    metrics["incomplete_case_nos"] = (
        df.loc[~df.pair_complete, "case_no"].astype(int).tolist()
    )

    gate = apply_locked_rule(metrics)
    csv_path = OUT_PRIV / "CALIBRATION_ROUND2_CODES.csv"
    df.to_csv(csv_path, index=False)

    payload = {
        "phase": "6C",
        "task": "calibration_round2_returns_ingested",
        "recorded_on": date.today().isoformat(),
        "codebook_version_at_coding": "0.3.0",
        "seed": 322720977,
        "blinded_packet_csv_sha256": sha256(PACKET),
        "returned_files": arch,
        "n_cases": 20,
        "n_complete_pairs": metrics["n_complete_pairs"],
        "n_incomplete_pairs": metrics["n_incomplete_pairs"],
        "incomplete_case_nos": metrics["incomplete_case_nos"],
        "agreement_complete_pairs": {
            "n_agree": metrics["n_agree_complete"],
            "n_disagree": metrics["n_disagree_complete"],
            "raw_agreement": metrics["raw_agreement"],
            "cohen_kappa_3way": metrics["cohen_kappa_3way"],
            "krippendorff_alpha_nominal": metrics["krippendorff_alpha_nominal"],
            "binary_explicit_raw_agreement": metrics["binary_explicit_raw_agreement"],
            "cohen_kappa_explicit_vs_rest": metrics["cohen_kappa_explicit_vs_rest"],
            "confusion_matrix_3way": metrics["confusion_matrix_3way"],
            "disagree_case_nos": metrics["disagree_case_nos_complete"],
        },
        "locked_rule_decision": gate,
        "cases": rows,
        "note": (
            "Round-2 agreement before discussion. Locked A/B/C rule applied. "
            "Not XIV main outcomes. Joint clarifications from Round-1 disagreement "
            "review feed v0.4 only if gate A or B."
        ),
    }
    (OUT_PRIV / "CALIBRATION_ROUND2_RETURNS.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    slim = {k: v for k, v in payload.items() if k != "cases"}
    REPORT_JSON.write_text(json.dumps(slim, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    write_agreement_md(payload)

    v04 = write_codebook_v04(gate["gate"])
    if v04 and JOINT_JSON.is_file():
        j = json.loads(JOINT_JSON.read_text(encoding="utf-8"))
        j["codebook_v04_written"] = True
        j["codebook_v04_gate"] = gate["gate"]
        JOINT_JSON.write_text(json.dumps(j, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    return {
        "metrics": metrics,
        "gate": gate,
        "csv": str(csv_path.relative_to(ROOT)),
        "report_md": str(REPORT_MD.relative_to(ROOT)),
        "codebook_v04": str(v04.relative_to(ROOT)) if v04 else None,
        **{k: v["path"] for k, v in arch.items()},
    }


def main() -> int:
    import argparse

    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--daniel", type=Path, required=True)
    p.add_argument("--jose-jaime", type=Path, required=True)
    p.add_argument("--joint", type=Path, default=None)
    args = p.parse_args()
    out = ingest(args.daniel, args.jose_jaime, args.joint)
    print(json.dumps({"gate": out["gate"], "metrics": out["metrics"], "v04": out["codebook_v04"]}, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
