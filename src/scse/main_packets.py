"""Generator and templates for future main-annotation DOCX packets.

Phase 6A: generator ready; human packets NOT issued.
Instructions remain provisional until codebook v0.4 is frozen.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from scse.main_frame import MAIN_ANNOTATION, MAIN_PRIVATE, VOLUMES_DEFAULT, load_main_corpus
from scse.paths import PROJECT_ROOT

RECIPIENTS = (
    ("Daniel", "Daniel Pinto Pajares"),
    ("Jose_Jaime", "Jose Jaime Baena Rojas"),
)

PLACEHOLDER_INTRO = (
    "Este cuaderno es una plantilla de trabajo preparada antes de congelar "
    "la guía. Todavía no es el paquete definitivo de anotación principal. "
    "No empecéis a codificar el corpus XIV hasta que recibáis la versión "
    "congelada tras el cierre de la calibración."
)


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.page_break_before = page_break_before
    set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def add_label(doc, label):
    return add_para(doc, label, size=12, bold=True, space_after=4, space_before=10)


def add_body_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Cm(0.2)
    set_run_font(p.add_run(str(text).strip()), size=11)
    return p


def add_checkbox_line(doc, label: str) -> None:
    add_para(doc, f"☐  {label}", size=12, space_after=2)


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _volume_slices(corpus: pd.DataFrame) -> list[pd.DataFrame]:
    plan_path = MAIN_PRIVATE / "MAIN_FRAME_VOLUME_PLAN.json"
    if plan_path.is_file():
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        slices = []
        for vol in plan["volumes"]:
            slices.append(corpus[corpus["case_no"].isin(vol["case_nos"])].copy())
        return slices
    # fallback contiguous
    out = []
    i = 0
    case_nos = corpus["case_no"].tolist()
    for size in VOLUMES_DEFAULT:
        chunk = case_nos[i : i + size]
        out.append(corpus[corpus["case_no"].isin(chunk)].copy())
        i += size
    return out


def build_one_volume_docx(
    *,
    recipient_slug: str,
    display_name: str,
    volume_df: pd.DataFrame,
    volume_no: int,
    n_volumes: int,
    issued: bool = False,
) -> Path:
    if issued:
        raise RuntimeError("Phase 6A forbids issuing final human packets")

    MAIN_ANNOTATION.mkdir(parents=True, exist_ok=True)
    out_dir = MAIN_ANNOTATION / "templates" / recipient_slug
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"main_{recipient_slug}_volume_{volume_no:02d}_TEMPLATE.docx"

    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        add_page_number(section)

    add_para(
        doc,
        f"Anotación principal — volumen {volume_no} de {n_volumes} (plantilla)",
        size=15,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Tipo de respuesta en las sesiones de control parlamentario",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_para(
        doc,
        f"Para {display_name}",
        size=12,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(doc, PLACEHOLDER_INTRO, space_after=10)
    add_para(
        doc,
        (
            f"Este volumen contiene {len(volume_df)} casos "
            f"(CASO {int(volume_df['case_no'].iloc[0])}–"
            f"{int(volume_df['case_no'].iloc[-1])}). "
            "El orden es el mismo para ambos anotadores."
        ),
        space_after=10,
    )
    add_para(
        doc,
        "¿Hasta qué punto la primera respuesta del Presidente del Gobierno responde a la cuestión planteada?",
        size=12,
        bold=True,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )
    add_para(doc, "Respuesta explícita", bold=True, space_after=2)
    add_para(doc, "La respuesta atiende directamente a la cuestión principal.", space_after=6)
    add_para(doc, "Respuesta parcial o intermedia", bold=True, space_after=2)
    add_para(
        doc,
        "Atiende una parte relevante de la cuestión, pero deja otra sin resolver o exige una inferencia razonable.",
        space_after=6,
    )
    add_para(doc, "Ausencia de respuesta", bold=True, space_after=2)
    add_para(
        doc,
        "Aunque hable del mismo tema, no aporta información que responda de hecho a la pregunta.",
        space_after=8,
    )
    add_para(
        doc,
        "Hablar del mismo tema no es necesariamente responder a la pregunta.",
        italic=True,
        space_after=10,
    )
    add_para(
        doc,
        "No hace falta justificar todos los casos. Usad observaciones solo cuando haya una duda real.",
        space_after=12,
    )

    for i, row in volume_df.iterrows():
        add_para(
            doc,
            f"CASO {int(row['case_no'])}",
            size=14,
            bold=True,
            space_after=6,
            page_break_before=True,
        )
        add_label(doc, "Pregunta registrada")
        add_body_block(doc, row["registered_question"])
        add_label(doc, "Pregunta oral")
        add_body_block(doc, row["q1_text"])
        add_label(doc, "Primera respuesta del Presidente del Gobierno")
        add_body_block(doc, row["r1_text"])
        add_label(doc, "Tu valoración")
        add_checkbox_line(doc, "Respuesta explícita")
        add_checkbox_line(doc, "Respuesta parcial o intermedia")
        add_checkbox_line(doc, "Ausencia de respuesta")
        add_label(doc, "¿Te ha parecido un caso dudoso?")
        add_checkbox_line(doc, "Sí")
        add_checkbox_line(doc, "No")
        add_label(doc, "Observaciones, si hacen falta:")
        add_para(doc, "_" * 72, space_after=2)
        add_para(doc, "_" * 72, space_after=8)

    doc.save(path)
    return path


def build_template_packets(*, issued: bool = False) -> dict:
    """Build volume templates for both recipients. Does not issue finals."""
    if issued:
        raise RuntimeError("Final human packets must not be issued in Phase 6A")
    corpus = load_main_corpus()
    slices = _volume_slices(corpus)
    n_vol = len(slices)
    paths: dict[str, list[str]] = {}
    for slug, name in RECIPIENTS:
        paths[slug] = []
        for vol_i, vol_df in enumerate(slices, start=1):
            p = build_one_volume_docx(
                recipient_slug=slug,
                display_name=name,
                volume_df=vol_df,
                volume_no=vol_i,
                n_volumes=n_vol,
                issued=False,
            )
            paths[slug].append(str(p.relative_to(PROJECT_ROOT)))

    # Determinism check: same case order for both
    d_cases = []
    j_cases = []
    for p in paths["Daniel"]:
        d_cases.extend(_extract_case_nos(PROJECT_ROOT / p))
    for p in paths["Jose_Jaime"]:
        j_cases.extend(_extract_case_nos(PROJECT_ROOT / p))
    if d_cases != j_cases:
        raise ValueError("Daniel/Jose Jaime template case orders differ")
    if d_cases != corpus["case_no"].tolist():
        raise ValueError("template case order differs from main corpus")

    meta = {
        "phase": "6A",
        "issued": False,
        "codebook_status": "awaiting_v0.4_after_round2",
        "n_cases": int(len(corpus)),
        "n_volumes": n_vol,
        "volume_sizes": [len(s) for s in slices],
        "paths": paths,
        "identical_case_order": True,
    }
    meta_path = MAIN_ANNOTATION / "MAIN_PACKET_GENERATOR_STATUS.json"
    meta_path.write_text(json.dumps(meta, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return meta


def _extract_case_nos(path: Path) -> list[int]:
    import re

    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        m = re.match(r"^CASO\s+(\d+)\s*$", p.text.strip())
        if m:
            out.append(int(m.group(1)))
    return out
