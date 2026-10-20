"""Build INTERNAL source-review DOCX/PDF for the pre-outcome main corpus.

No reply-status checkboxes. No AI interpretation.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from scse.main_frame import MAIN_PACKETS, load_main_corpus
from scse.paths import PROJECT_ROOT


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.page_break_before = page_break_before
    set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def add_label(doc, label):
    return add_para(doc, label, size=12, bold=True, space_after=4, space_before=10)


def add_body_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Cm(0.2)
    set_run_font(p.add_run(str(text).strip()), size=11)
    return p


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def build_source_review_docx(path: Path | None = None) -> Path:
    corpus = load_main_corpus()
    MAIN_PACKETS.mkdir(parents=True, exist_ok=True)
    path = path or (MAIN_PACKETS / "MAIN_CORPUS_SOURCE_REVIEW.docx")

    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        add_page_number(section)

    add_para(
        doc,
        "Revisión de fuentes del corpus principal",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Tipo de respuesta en las sesiones de control parlamentario",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(
        doc,
        (
            "Documento interno de comprobación de extracción y procedencia. "
            f"Incluye los {len(corpus)} intercambios del marco principal "
            "pre-anotación (XIV; sin sondas de Fase 0). "
            "No es un cuaderno de anotación: no pide valoración de respuesta."
        ),
        space_after=10,
    )
    add_para(
        doc,
        (
            "Para cada caso se muestra la pregunta registrada, la pregunta oral "
            "y la primera respuesta del Presidente del Gobierno, tal como "
            "quedan en el corpus privado."
        ),
        space_after=14,
    )

    for i, row in corpus.iterrows():
        first = i == 0
        add_para(
            doc,
            f"CASO {int(row['case_no'])}",
            size=14,
            bold=True,
            space_before=6 if not first else 0,
            space_after=6,
            page_break_before=not first,
        )
        add_para(
            doc,
            f"Unidad: {row['exchange_id']}  ·  {row['expediente']}  ·  {row['date']}",
            size=10,
            italic=True,
            space_after=8,
        )
        add_label(doc, "Pregunta registrada")
        add_body_block(doc, row["registered_question"])
        add_label(doc, "Pregunta oral")
        add_body_block(doc, row["q1_text"])
        add_label(doc, "Primera respuesta del Presidente del Gobierno")
        add_body_block(doc, row["r1_text"])

    doc.save(path)
    return path


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path | None = None) -> Path:
    import subprocess

    pdf_path = pdf_path or docx_path.with_suffix(".pdf")
    out_dir = pdf_path.parent
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
    )
    produced = out_dir / (docx_path.stem + ".pdf")
    if produced.resolve() != pdf_path.resolve():
        produced.replace(pdf_path)
    return pdf_path


def build_source_review_pdf() -> dict:
    docx = build_source_review_docx()
    pdf = convert_docx_to_pdf(docx, MAIN_PACKETS / "MAIN_CORPUS_SOURCE_REVIEW.pdf")
    return {
        "docx": str(docx.relative_to(PROJECT_ROOT)),
        "pdf": str(pdf.relative_to(PROJECT_ROOT)),
        "n_cases": len(load_main_corpus()),
    }
