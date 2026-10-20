"""Private QA30 linkage-review viewer (project-lead inspection only).

No reply_status. No political alignment. No outcome interpretation.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from scse.main_frame import load_qa30_exchange_ids
from scse.paths import DATA_PRIVATE, LINK_AUDIT_PRIVATE, PROJECT_ROOT

POOL = DATA_PRIVATE / "eligible_pool_full.parquet"
QA_CSV = LINK_AUDIT_PRIVATE / "PHASE2_LINKAGE_QA_HUMAN_VALIDATED.csv"
QA_JSON = LINK_AUDIT_PRIVATE / "PHASE2_LINKAGE_QA_HUMAN_VALIDATION.json"
OUT_DOCX = LINK_AUDIT_PRIVATE / "QA30_LINKAGE_REVIEW_VIEW.docx"
OUT_PDF = LINK_AUDIT_PRIVATE / "QA30_LINKAGE_REVIEW_VIEW.pdf"


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text, *, size=12, bold=False, italic=False, space_after=8, align=None, page_break_before=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.page_break_before = page_break_before
    set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def build_qa30_viewer() -> dict:
    qa_ids = load_qa30_exchange_ids()
    meta = json.loads(QA_JSON.read_text(encoding="utf-8"))
    validated = pd.read_csv(QA_CSV).set_index("exchange_id")
    pool = pd.read_parquet(POOL).set_index("exchange_id")

    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        add_page_number(section)

    add_para(
        doc,
        "QA30 linkage review viewer (private)",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Project-lead inspection only. Linkage provenance — not reply-status coding.",
        size=11,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(
        doc,
        (
            f"Reviewer role on record: {meta.get('reviewer_type')} "
            f"(exact personal name not stored in the validation JSON). "
            f"Modality: {meta.get('review_modality')}. "
            f"reply_status_judgement_performed={meta.get('reply_status_judgement_performed')}. "
            "Daniel Pinto Pajares and Jose Jaime Baena Rojas are not named in QA provenance files."
        ),
        space_after=14,
    )

    for i, eid in enumerate(qa_ids, start=1):
        row = pool.loc[eid]
        v = validated.loc[eid]
        add_para(
            doc,
            f"QA CASO {i}",
            size=14,
            bold=True,
            page_break_before=(i > 1),
            space_after=4,
        )
        add_para(
            doc,
            f"exchange_id: {eid}  ·  {row.get('expediente')}  ·  {row.get('date')}",
            size=10,
            italic=True,
            space_after=8,
        )
        add_para(doc, "Pregunta registrada", bold=True, space_after=2)
        add_para(doc, str(row["registered_question"]).strip(), size=11, space_after=8)
        add_para(doc, "Pregunta oral", bold=True, space_after=2)
        add_para(doc, str(row["q1_text"]).strip(), size=11, space_after=8)
        add_para(doc, "Primera respuesta del Presidente del Gobierno", bold=True, space_after=2)
        add_para(doc, str(row["r1_text"]).strip(), size=11, space_after=8)
        add_para(doc, "Linkage judgement", bold=True, space_after=2)
        note = v.get("human_notes")
        note_s = "" if pd.isna(note) or str(note).strip() in {"", "nan"} else str(note)
        add_para(
            doc,
            (
                f"human_link_correct={v.get('human_link_correct')}; "
                f"link_status={row.get('link_status')}; "
                f"source={v.get('human_validation_source')}"
                + (f"; note={note_s}" if note_s else "; note=(none recorded)")
            ),
            size=11,
            space_after=6,
        )
        add_para(
            doc,
            (
                "Provenance status: oral aggregate YES recorded at Phase-2.5 gate; "
                "original spreadsheet cells were never filled."
            ),
            size=10,
            italic=True,
            space_after=8,
        )

    LINK_AUDIT_PRIVATE.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)

    import subprocess

    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(LINK_AUDIT_PRIVATE),
            str(OUT_DOCX),
        ],
        check=True,
        capture_output=True,
    )
    return {
        "docx": str(OUT_DOCX.relative_to(PROJECT_ROOT)),
        "pdf": str(OUT_PDF.relative_to(PROJECT_ROOT)),
        "n": 30,
        "reply_status_included": False,
    }
