"""Human Word packets for Calibration Round 1 (Daniel / Jose Jaime).

Polish pass only: same 20 cases, same order, same source texts.
Natural academic Spanish. No technical metadata on the page.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from scse.paths import DATA_PRIVATE, PROJECT_ROOT, PROTOCOL_DIR

ROOT = PROJECT_ROOT
PACKET_CSV = (
    DATA_PRIVATE / "calibration" / "round1" / "CALIBRATION_ROUND1_PACKET_BLINDED.csv"
)
OUT_DIR = ROOT / "_internal" / "calibration_round1"
CODEBOOK = PROTOCOL_DIR / "reply_status_codebook_v0.3.md"
MANIFEST = PROTOCOL_DIR / "calibration_round1_manifest.yaml"


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_label(doc, label):
    return add_para(
        doc,
        label,
        size=12,
        bold=True,
        space_after=4,
        space_before=10,
    )


def add_body_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Cm(0.2)
    set_run_font(p.add_run(str(text).strip()), size=11)
    return p


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "12"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "666666"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def add_response_block(doc, *, blank_lines: int = 3) -> None:
    """Valuation + dudoso + writing space in one unsplittable block."""
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_border(cell, val="nil", sz="0", color="FFFFFF")
    cell.text = ""
    first = True

    def _line(text: str, *, bold=False, after=4, before=0, size=12):
        nonlocal first
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        set_run_font(p.add_run(text), size=size, bold=bold)

    _line("Tu valoración", bold=True, after=6, before=4)
    _line("☐  Respuesta explícita", after=3)
    _line("☐  Respuesta parcial o intermedia", after=3)
    _line("☐  Ausencia de respuesta", after=8)
    _line("¿Te ha parecido un caso dudoso?", bold=True, after=4)
    _line("☐  Sí", after=2)
    _line("☐  No", after=8)
    _line("Observaciones, si hacen falta:", bold=True, after=2)

    # Writing area: same cell, so it cannot orphan onto the next page alone.
    rule = cell.add_paragraph()
    rule.paragraph_format.space_after = Pt(0)
    rule.paragraph_format.space_before = Pt(2)
    set_run_font(rule.add_run("_" * 64), size=10)

    for _ in range(blank_lines):
        blank = cell.add_paragraph()
        blank.paragraph_format.space_after = Pt(10)
        set_run_font(blank.add_run(" "), size=11)

    foot = cell.add_paragraph()
    foot.paragraph_format.space_before = Pt(0)
    foot.paragraph_format.space_after = Pt(2)
    set_run_font(foot.add_run("_" * 64), size=10)

    add_para(doc, "", size=4, space_after=2)


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10)
    # PAGE field
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)


def clear_core_props(doc: Document, *, recipient: str) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = f"Calibración — primera ronda ({recipient})"
    props.subject = ""
    props.keywords = ""
    props.category = ""
    props.comments = ""


def add_instructions(doc: Document, *, recipient: str) -> None:
    doc.add_paragraph()
    add_para(
        doc,
        "Primera ronda de calibración",
        size=20,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Tipo de respuesta en las sesiones de control parlamentario",
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_para(
        doc,
        f"Para {recipient}",
        size=12,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=16,
    )

    add_para(
        doc,
        (
            "Os paso una nueva tanda de 20 casos. Esta vez es importante que "
            "los hagáis cada uno por vuestra cuenta."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "Haced esta ronda por separado y no comentéis los casos entre "
            "vosotros hasta que me hayáis devuelto los dos documentos."
        ),
        size=12,
        space_after=12,
    )

    add_para(
        doc,
        (
            "En cada caso hay que valorar hasta qué punto la primera respuesta "
            "del Presidente responde a la cuestión planteada."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "¿Hasta qué punto la primera respuesta del Presidente del Gobierno "
            "responde a la cuestión planteada?"
        ),
        size=12,
        bold=True,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )

    add_para(doc, "Respuesta explícita", size=12, bold=True, space_after=2)
    add_para(
        doc,
        "La respuesta atiende directamente a la cuestión principal.",
        size=12,
        space_after=8,
    )
    add_para(doc, "Respuesta parcial o intermedia", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "Responde a una parte relevante de la cuestión, pero deja otra "
            "parte sin resolver o exige una inferencia razonable."
        ),
        size=12,
        space_after=8,
    )
    add_para(doc, "Ausencia de respuesta", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "Aunque pueda hablar del mismo tema, no aporta información que "
            "responda realmente a la cuestión planteada."
        ),
        size=12,
        space_after=10,
    )
    add_para(
        doc,
        "Hablar del mismo tema no es necesariamente responder a la pregunta.",
        size=12,
        italic=True,
        space_after=10,
    )
    add_para(
        doc,
        (
            "No se valora si la respuesta es verdadera, si resulta convincente "
            "ni si políticamente se está de acuerdo con ella. La pregunta oral "
            "es el objetivo principal; la pregunta registrada sirve de contexto."
        ),
        size=12,
        space_after=10,
    )
    add_para(
        doc,
        (
            "No hace falta justificar todos los casos. Utilizad el espacio de "
            "observaciones solo cuando tengáis dudas o veáis algún problema "
            "con el criterio."
        ),
        size=12,
        space_after=6,
    )


def add_case(
    doc: Document,
    case_no: int,
    *,
    registered: str,
    q1: str,
    r1: str,
) -> None:
    add_para(
        doc,
        f"CASO {case_no}",
        size=16,
        bold=True,
        space_after=8,
        space_before=0,
        page_break_before=True,
    )

    add_label(doc, "Pregunta registrada")
    add_body_block(doc, registered)

    add_label(doc, "Pregunta oral")
    add_body_block(doc, q1)

    add_label(doc, "Primera respuesta del Presidente del Gobierno")
    add_body_block(doc, r1)

    add_response_block(doc)


def add_closing(doc: Document) -> None:
    # Continue after the last case; avoid a nearly empty final page when possible.
    add_para(
        doc,
        (
            "Si algún caso os ha generado una duda que no encaja bien con las "
            "tres opciones, dejadlo indicado en observaciones. Nos interesa "
            "especialmente detectar esos casos antes de la siguiente ronda."
        ),
        size=12,
        italic=True,
        space_after=8,
        space_before=14,
    )
    add_para(
        doc,
        "Gracias por el tiempo y por la lectura cuidadosa.",
        size=12,
        italic=True,
        space_after=4,
    )


def build_one(annotator_slug: str, display_name: str) -> Path:
    df = pd.read_csv(PACKET_CSV)
    assert len(df) == 20
    assert list(df.sort_values("case_no")["case_no"]) == list(range(1, 21))

    doc = Document()
    clear_core_props(doc, recipient=display_name)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_page_number(section)

    add_instructions(doc, recipient=display_name)

    rows = df.sort_values("case_no").to_dict(orient="records")
    for row in rows:
        add_case(
            doc,
            int(row["case_no"]),
            registered=str(row["registered_question"]),
            q1=str(row["Q1"]),
            r1=str(row["R1"]),
        )

    add_closing(doc)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"calibracion_ronda1_{annotator_slug}.docx"
    doc.save(out)
    return out


def write_provenance(paths: dict[str, Path]) -> Path:
    payload = {
        "phase": "5A",
        "round": 1,
        "task": "human_facing_docx_polish",
        "scientific_draw_changed": False,
        "codebook_version": "0.3.0",
        "codebook_sha256": hashlib.sha256(CODEBOOK.read_bytes()).hexdigest(),
        "manifest_yaml": "zenodo/protocol/calibration_round1_manifest.yaml",
        "manifest_yaml_sha256": hashlib.sha256(MANIFEST.read_bytes()).hexdigest(),
        "blinded_packet_csv": str(PACKET_CSV.relative_to(ROOT)),
        "blinded_packet_csv_sha256": hashlib.sha256(PACKET_CSV.read_bytes()).hexdigest(),
        "packets": {
            name: {
                "path": str(path.relative_to(ROOT)),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
            for name, path in paths.items()
        },
        "same_cases_same_order": True,
        "self_contained_word": True,
        "independence_rule": (
            "No discussion until both completed packets are returned; "
            "agreement calculated before discussion."
        ),
    }
    out = OUT_DIR / "CALIBRATION_ROUND1_PACKET_PROVENANCE.json"
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return out


def main() -> int:
    daniel = build_one("Daniel", "Daniel Pinto Pajares")
    jose = build_one("Jose_Jaime", "Jose Jaime Baena Rojas")
    prov = write_provenance({"daniel": daniel, "jose_jaime": jose})
    print(daniel)
    print(jose)
    print(prov)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
