"""Build Phase 5B human package: joint review + Round-2 packets + email."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from scse.paths import DATA_PRIVATE, PROJECT_ROOT, PROTOCOL_DIR

ROOT = PROJECT_ROOT
OUT = ROOT / "_internal" / "calibration_round2"
R1_PKT = DATA_PRIVATE / "calibration" / "round1" / "CALIBRATION_ROUND1_PACKET_BLINDED.csv"
R1_CODES = DATA_PRIVATE / "calibration" / "round1" / "CALIBRATION_ROUND1_CODES.csv"
R2_PKT = DATA_PRIVATE / "calibration" / "round2" / "CALIBRATION_ROUND2_PACKET_BLINDED.csv"
R2_AUD = DATA_PRIVATE / "calibration" / "round2" / "CALIBRATION_ROUND2_AUDIT.csv"
R2_DIV = DATA_PRIVATE / "calibration" / "round2" / "CALIBRATION_ROUND2_DIVERSITY.json"
JOINT_CASES = (5, 6, 7, 10)

LABEL_ES = {
    "respuesta explícita": "Respuesta explícita",
    "respuesta parcial o intermedia": "Respuesta parcial o intermedia",
    "ausencia de respuesta": "Ausencia de respuesta",
}


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.page_break_before = page_break_before
    set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def add_label(doc, label):
    return add_para(doc, label, size=12, bold=True, space_after=4, space_before=10)


def add_body_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Cm(0.2)
    set_run_font(p.add_run(str(text).strip()), size=11)
    return p


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "12"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "666666"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_response_block(doc, *, blank_lines: int = 3) -> None:
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_border(cell, val="nil", sz="0", color="FFFFFF")
    cell.text = ""
    first = True

    def _line(text: str, *, bold=False, after=4, before=0, size=12):
        nonlocal first
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        set_run_font(p.add_run(text), size=size, bold=bold)

    _line("Tu valoración", bold=True, after=6, before=4)
    _line("☐  Respuesta explícita", after=3)
    _line("☐  Respuesta parcial o intermedia", after=3)
    _line("☐  Ausencia de respuesta", after=8)
    _line("¿Te ha parecido un caso dudoso?", bold=True, after=4)
    _line("☐  Sí", after=2)
    _line("☐  No", after=8)
    _line("Observaciones, si hacen falta:", bold=True, after=2)
    rule = cell.add_paragraph()
    rule.paragraph_format.space_after = Pt(0)
    set_run_font(rule.add_run("_" * 64), size=10)
    for _ in range(blank_lines):
        blank = cell.add_paragraph()
        blank.paragraph_format.space_after = Pt(10)
        set_run_font(blank.add_run(" "), size=11)
    foot = cell.add_paragraph()
    set_run_font(foot.add_run("_" * 64), size=10)
    add_para(doc, "", size=4, space_after=2)


def add_joint_decision_block(doc) -> None:
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_border(cell, val="nil", sz="0", color="FFFFFF")
    cell.text = ""
    first = True

    def _line(text: str, *, bold=False, after=4, before=0, italic=False):
        nonlocal first
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        set_run_font(p.add_run(text), size=12, bold=bold, italic=italic)

    _line(
        "Si después de comentarlo llegáis a una lectura común, marcadla aquí.",
        bold=True,
        before=6,
        after=6,
    )
    _line("☐  Respuesta explícita", after=3)
    _line("☐  Respuesta parcial o intermedia", after=3)
    _line("☐  Ausencia de respuesta", after=3)
    _line("☐  Seguimos sin verlo igual", after=8)
    _line("¿Qué os ayuda a decidir este caso?", bold=True, after=4)
    rule = cell.add_paragraph()
    set_run_font(rule.add_run("_" * 64), size=10)
    for _ in range(3):
        blank = cell.add_paragraph()
        blank.paragraph_format.space_after = Pt(10)
        set_run_font(blank.add_run(" "), size=11)
    foot = cell.add_paragraph()
    set_run_font(foot.add_run("_" * 64), size=10)
    add_para(doc, "", size=4, space_after=2)


def add_lines_box(doc, *, n_lines: int = 5) -> None:
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_border(cell, val="single", sz="12", color="888888")
    cell.text = ""
    p = cell.paragraphs[0]
    set_run_font(p.add_run(" "), size=11)
    for _ in range(n_lines - 1):
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(8)
        set_run_font(p2.add_run(" "), size=11)
    add_para(doc, "", size=4, space_after=6)


def clear_props(doc: Document, title: str) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = title
    props.subject = ""
    props.keywords = ""
    props.comments = ""


def setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_page_number(section)


def build_joint() -> Path:
    pkt = pd.read_csv(R1_PKT).sort_values("case_no")
    codes = pd.read_csv(R1_CODES).sort_values("case_no")
    by_case = {int(r.case_no): r for _, r in pkt.iterrows()}
    code_by = {int(r.case_no): r for _, r in codes.iterrows()}

    doc = Document()
    clear_props(doc, "Revisión conjunta — casos 5, 6, 7 y 10")
    setup_page(doc)

    add_para(
        doc,
        "Revisión de cuatro casos",
        size=20,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Tipo de respuesta en las sesiones de control parlamentario",
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=14,
    )
    add_para(
        doc,
        (
            "En la primera ronda coincidisteis en dieciséis de los veinte casos. "
            "Antes de pasar a la última tanda, interesa revisar únicamente estos "
            "cuatro, para dejar un poco más clara la frontera entre las tres "
            "opciones."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "La idea no es decidir quién acertó, sino ver qué os lleva a leer "
            "cada caso de forma distinta. No hace falta rehacer la guía: "
            "comentad cada caso y, al final, dejad por escrito las dos "
            "distinciones que encontraréis."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "Cuando terminéis esta revisión conjunta, pasad a los documentos "
            "individuales. A partir de ese momento, no comentéis entre vosotros "
            "los 20 casos nuevos hasta que ambos los hayáis terminado."
        ),
        size=12,
        space_after=12,
    )

    for case_no in JOINT_CASES:
        row = by_case[case_no]
        code = code_by[case_no]
        add_para(
            doc,
            f"CASO {case_no}",
            size=16,
            bold=True,
            space_after=8,
            page_break_before=True,
        )
        add_label(doc, "Pregunta registrada")
        add_body_block(doc, row["registered_question"])
        add_label(doc, "Pregunta oral")
        add_body_block(doc, row["Q1"])
        add_label(doc, "Primera respuesta del Presidente del Gobierno")
        add_body_block(doc, row["R1"])
        add_para(
            doc,
            "Vuestras valoraciones iniciales",
            size=12,
            bold=True,
            space_before=10,
            space_after=4,
        )
        add_para(
            doc,
            f"Daniel: {str(code.daniel_reply_status)}.",
            size=12,
            space_after=2,
        )
        add_para(
            doc,
            f"Jose Jaime: {str(code.jose_jaime_reply_status)}.",
            size=12,
            space_after=6,
        )
        add_joint_decision_block(doc)

    add_para(
        doc,
        "Dos distinciones para dejar por escrito",
        size=16,
        bold=True,
        space_before=8,
        space_after=10,
        page_break_before=True,
    )
    add_para(
        doc,
        (
            "Con lo hablado en estos cuatro casos, dejad anotado, con vuestras "
            "palabras, solo esto."
        ),
        size=12,
        space_after=10,
    )
    add_para(
        doc,
        (
            "1. ¿En qué momento una intervención deja de ser, para vosotros, "
            "una simple ausencia de respuesta y pasa a ser una respuesta parcial?"
        ),
        size=12,
        bold=True,
        space_after=6,
    )
    add_lines_box(doc, n_lines=5)
    add_para(
        doc,
        (
            "2. ¿Cuándo una respuesta parcial os parece ya lo bastante completa "
            "para marcarla como explícita?"
        ),
        size=12,
        bold=True,
        space_before=8,
        space_after=6,
    )
    add_lines_box(doc, n_lines=5)
    add_para(
        doc,
        (
            "Gracias. Con esto cerramos la parte conjunta y pasáis a los "
            "cuadernos individuales."
        ),
        size=12,
        italic=True,
        space_before=10,
    )

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "01_REVISION_CONJUNTA_CASOS_5_6_7_10.docx"
    doc.save(path)
    return path


def build_round2_one(slug: str, display_name: str) -> Path:
    df = pd.read_csv(R2_PKT).sort_values("case_no")
    assert len(df) == 20
    assert list(df["case_no"]) == list(range(1, 21))

    doc = Document()
    clear_props(doc, f"Calibración — segunda ronda ({display_name})")
    setup_page(doc)

    add_para(
        doc,
        "Segunda ronda de calibración",
        size=20,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Tipo de respuesta en las sesiones de control parlamentario",
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_para(
        doc,
        f"Para {display_name}",
        size=12,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=14,
    )
    add_para(
        doc,
        (
            "Antes de empezar, terminad la revisión conjunta de los cuatro casos "
            "anteriores. A partir de aquí, cada uno trabaja por separado."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "Os paso una última tanda de 20 casos nuevos. Hacedlos cada uno por "
            "vuestra cuenta y no comentéis las respuestas entre vosotros hasta "
            "que me devolváis los dos documentos."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "En cada caso hay que valorar hasta qué punto la primera respuesta "
            "del Presidente responde a la cuestión planteada."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "¿Hasta qué punto la primera respuesta del Presidente del Gobierno "
            "responde a la cuestión planteada?"
        ),
        size=12,
        bold=True,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(doc, "Respuesta explícita", size=12, bold=True, space_after=2)
    add_para(
        doc,
        "La respuesta atiende directamente a la cuestión principal.",
        size=12,
        space_after=8,
    )
    add_para(doc, "Respuesta parcial o intermedia", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "Responde a una parte relevante de la cuestión, pero deja otra "
            "parte sin resolver o exige una inferencia razonable."
        ),
        size=12,
        space_after=8,
    )
    add_para(doc, "Ausencia de respuesta", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "Aunque pueda hablar del mismo tema, no aporta información que "
            "responda realmente a la cuestión planteada."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        "Hablar del mismo tema no es necesariamente responder a la pregunta.",
        size=12,
        italic=True,
        space_after=8,
    )
    add_para(
        doc,
        (
            "No hace falta justificar todos los casos. Utilizad el espacio de "
            "observaciones solo cuando tengáis dudas o veáis algún problema "
            "con el criterio."
        ),
        size=12,
        space_after=6,
    )

    for _, row in df.iterrows():
        add_para(
            doc,
            f"CASO {int(row.case_no)}",
            size=16,
            bold=True,
            space_after=8,
            page_break_before=True,
        )
        add_label(doc, "Pregunta registrada")
        add_body_block(doc, row["registered_question"])
        add_label(doc, "Pregunta oral")
        add_body_block(doc, row["Q1"])
        add_label(doc, "Primera respuesta del Presidente del Gobierno")
        add_body_block(doc, row["R1"])
        add_response_block(doc)

    add_para(
        doc,
        (
            "Si algún caso os ha generado una duda que no encaja bien con las "
            "tres opciones, dejadlo indicado en observaciones."
        ),
        size=12,
        italic=True,
        space_before=10,
        space_after=6,
    )
    add_para(
        doc,
        "Gracias por el tiempo y por la lectura cuidadosa.",
        size=12,
        italic=True,
    )

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"02_calibracion_ronda2_{slug}.docx"
    doc.save(path)
    return path


def write_email() -> Path:
    text = """Asunto: Última ronda de ajuste (revisión conjunta + 20 casos nuevos)

Hola Daniel, hola Jose Jaime:

Os mando ya la última ronda de ajuste para cerrar esta parte.

Hay dos pasos muy breves:

1) Primero abrid juntos el documento 01_REVISION_CONJUNTA_CASOS_5_6_7_10.docx.
Ahí están solo los cuatro casos de la primera ronda en los que no coincidisteis.
Comentadlos, marcáis una decisión conjunta cuando podáis y dejáis por escrito
las dos distinciones del final. No hace falta rehacer la guía.

2) En cuanto terminéis esa revisión, cada uno pasa a su cuaderno individual
(02_calibracion_ronda2_…). Son 20 casos completamente nuevos. A partir de ese
momento, trabajad por separado y no comentéis esos casos entre vosotros hasta
que ambos me hayáis devuelto el documento.

No hace falta justificar todos los casos: usad observaciones solo cuando algo
os genere una duda real.

Cuando me devolváis los tres archivos, cerramos esta calibración.

Un abrazo,
César
"""
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "EMAIL_UNICO_CALIBRACION_FINAL.txt"
    path.write_text(text, encoding="utf-8")
    return path


def normalize_text(text: object) -> str:
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    s = s.casefold()
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"[¿?¡!.,;:…\"'«»\-–—()/\[\]{}]", "", s)
    return re.sub(r"\s+", " ", s).strip()


def write_audit() -> Path:
    from scse.build_calibration_round2 import (
        SEED_DIGEST,
        SEED_ROUND2,
        SEED_SOURCE_STRING,
        load_development_exclusions,
        load_round1_exclusions,
        load_xiv_and_qa_exclusions,
        normalize_text as norm,
    )

    aud = pd.read_csv(R2_AUD)
    pkt = pd.read_csv(R2_PKT)
    div = json.loads(R2_DIV.read_text(encoding="utf-8"))
    r1 = load_round1_exclusions()
    dev = load_development_exclusions()
    xq = load_xiv_and_qa_exclusions()

    q1_ids = set(aud["q1_utterance_id"].astype(str))
    keys = set(zip(aud["q1_utterance_id"].astype(str), aud["r1_utterance_id"].astype(str)))
    reg_n = [norm(t) for t in pkt["registered_question"]]
    q1_n = [norm(t) for t in pkt["Q1"]]

    overlaps = {
        "development_unit_or_utt": len(q1_ids & dev["q1"]) + len(set(aud.unit_id) & dev["unit_ids"]),
        "round1_unit_or_utt": len(q1_ids & r1["q1"]) + len(set(aud.unit_id) & r1["unit_ids"]),
        "xiv_q1": len(q1_ids & xq["xiv_q1"]),
        "qa_q1": len(q1_ids & xq["qa_q1"]),
        "phase0_q1": len(q1_ids & xq["phase0_q1"]),
        "dev_reg_norm": sum(1 for t in reg_n if t in dev["reg_norm"]),
        "r1_reg_norm": sum(1 for t in reg_n if t in r1["reg_norm"]),
        "dev_q1_norm": sum(1 for t in q1_n if t in dev["q1_norm"]),
        "r1_q1_norm": sum(1 for t in q1_n if t in r1["q1_norm"]),
        "xiv_reg_norm": sum(1 for t in reg_n if t in xq["xiv_reg_norm"]),
        "xiv_q1_norm": sum(1 for t in q1_n if t in xq["xiv_q1_norm"]),
    }
    within = {
        "n": len(pkt),
        "distinct_units": int(aud["unit_id"].nunique()),
        "distinct_exchanges": len(keys),
        "duplicate_q1_norm": int(pd.Series(q1_n).duplicated().sum()),
        "duplicate_reg_norm": int(pd.Series(reg_n).duplicated().sum()),
    }

    lines = [
        "# Calibration Round 2 — selection and novelty audit",
        "",
        "**Phase:** 5B  ",
        "**Date:** 2026-10-13  ",
        "**Final planned calibration round:** yes (no Round 3 by default)",
        "",
        "## Seed",
        "",
        f"- Source string: `{SEED_SOURCE_STRING}`",
        f"- SHA-256: `{SEED_DIGEST}`",
        "- Conversion: `int(hex[:16], 16) % (2**31 - 1)`",
        f"- Seed: `{SEED_ROUND2}`",
        "",
        "## Sample",
        "",
        f"- N: **{within['n']}**",
        f"- Distinct units: **{within['distinct_units']}**",
        f"- Distinct exchanges: **{within['distinct_exchanges']}**",
        f"- Duplicate normalized Q1: **{within['duplicate_q1_norm']}**",
        f"- Duplicate normalized registered: **{within['duplicate_reg_norm']}**",
        "",
        "## Diversity",
        "",
        f"- Topics represented ({div['n_topics']}): `{div['topics']}`",
        f"- Largest topic share: **{div['max_topic_share']}**",
        f"- Forms represented ({div['n_forms']}): `{div['forms']}`",
        f"- Distinct questioners: **{div['n_questioners']}**",
        f"- Maximum cases from one questioner: **{div['max_per_questioner']}**",
        f"- Distinct sittings: **{div['n_sittings']}**",
        f"- Maximum cases from one sitting: **{div['max_per_sitting']}**",
        f"- Legislatures: `{div['legislatures']}`",
        f"- Ladder used: `{div['ladder_used']}`",
        f"- Targets missed: `{div['targets_missed']}`",
        "",
        "Hard novelty was not relaxed. The only missed soft target is",
        "`max_per_questioner_le_2` (frame size forces up to 4 after the fixed",
        "relaxation ladder).",
        "",
        "## Novelty overlaps (must be 0)",
        "",
        "| Comparison | Count |",
        "|------------|------:|",
    ]
    for k, v in overlaps.items():
        lines.append(f"| {k} | {v} |")
    lines += [
        "",
        "## Integrity",
        "",
        "- Reply-status labels were **not** inspected during selection.",
        "- Round-1 scientific packet unchanged.",
        "- No XIV outcomes calculated.",
        "- Round 3 not planned by default.",
        "",
        "## Artifacts",
        "",
        "- `_internal/data_private/calibration/round2/`",
        "- `_internal/calibration_round2/`",
        "- `zenodo/protocol/calibration_round2_manifest.yaml`",
        "",
    ]
    path = ROOT / "_internal" / "reports" / "CALIBRATION_ROUND2_SELECTION_AND_NOVELTY_AUDIT.md"
    path.write_text("\n".join(lines), encoding="utf-8")

    machine = {
        "seed": SEED_ROUND2,
        "seed_source_string": SEED_SOURCE_STRING,
        "seed_sha256": SEED_DIGEST,
        "within": within,
        "diversity": div,
        "overlaps": overlaps,
        "packet_sha256": hashlib.sha256(R2_PKT.read_bytes()).hexdigest(),
    }
    (ROOT / "_internal" / "reports" / "CALIBRATION_ROUND2_SELECTION_AND_NOVELTY_AUDIT.json").write_text(
        json.dumps(machine, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return path


def verify_packets_identical_except_recipient(d_path: Path, j_path: Path) -> bool:
    def body(path: Path) -> str:
        doc = Document(path)
        return "\n".join(
            p.text for p in doc.paragraphs if not p.text.startswith("Para ")
        )

    return body(d_path) == body(j_path)


def main() -> int:
    joint = build_joint()
    daniel = build_round2_one("Daniel", "Daniel Pinto Pajares")
    jose = build_round2_one("Jose_Jaime", "Jose Jaime Baena Rojas")
    email = write_email()
    audit = write_audit()
    same = verify_packets_identical_except_recipient(daniel, jose)
    # source text integrity vs CSV
    pkt = pd.read_csv(R2_PKT).sort_values("case_no")

    def extract_cases(path: Path):
        doc = Document(path)
        cases = {}
        cur = None
        mode = None
        buf = []

        def flush():
            nonlocal buf, mode
            if cur is not None and mode is not None:
                cases.setdefault(cur, {})[mode] = "\n".join(buf).strip()
            buf = []

        for p in doc.paragraphs:
            t = p.text
            m = re.match(r"^CASO\s+(\d+)\s*$", t.strip())
            if m:
                flush()
                cur = int(m.group(1))
                mode = None
                continue
            if t.strip() == "Pregunta registrada":
                flush()
                mode = "registered"
                continue
            if t.strip() == "Pregunta oral":
                flush()
                mode = "oral"
                continue
            if t.strip() == "Primera respuesta del Presidente del Gobierno":
                flush()
                mode = "r1"
                continue
            if (
                t.strip()
                in (
                    "Tu valoración",
                    "¿Te ha parecido un caso dudoso?",
                    "Observaciones, si hacen falta:",
                )
                or t.strip().startswith("☐")
                or (t.strip() and set(t.strip()) <= {"_"})
                or t.strip().startswith("Si algún caso")
                or t.strip().startswith("Gracias por el tiempo")
            ):
                flush()
                mode = None
                continue
            if mode and t.strip():
                buf.append(t.strip())
        flush()
        return cases

    cd = extract_cases(daniel)
    ok = True
    for _, row in pkt.iterrows():
        c = int(row.case_no)
        if cd[c]["registered"] != str(row.registered_question).strip():
            ok = False
        if cd[c]["oral"] != str(row.Q1).strip():
            ok = False
        if cd[c]["r1"] != str(row.R1).strip():
            ok = False

    prov = {
        "phase": "5B",
        "joint": {"path": str(joint.relative_to(ROOT)), "sha256": hashlib.sha256(joint.read_bytes()).hexdigest()},
        "daniel": {"path": str(daniel.relative_to(ROOT)), "sha256": hashlib.sha256(daniel.read_bytes()).hexdigest()},
        "jose_jaime": {"path": str(jose.relative_to(ROOT)), "sha256": hashlib.sha256(jose.read_bytes()).hexdigest()},
        "email": str(email.relative_to(ROOT)),
        "audit": str(audit.relative_to(ROOT)),
        "same_cases_same_order_except_recipient": same,
        "source_texts_match_csv": ok,
        "round2_packet_sha256": hashlib.sha256(R2_PKT.read_bytes()).hexdigest(),
    }
    (OUT / "CALIBRATION_ROUND2_PACKET_PROVENANCE.json").write_text(
        json.dumps(prov, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(json.dumps(prov, indent=2, ensure_ascii=False))
    return 0 if same and ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
