"""Build human-friendly development annotation DOCX (Phase 4.2)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "_internal" / "development" / "development_annotator_A.xlsx"
OUT_DOCX = ROOT / "_internal" / "development" / "Reply_Status_Development_Exercise.docx"
OUT_PDF = ROOT / "_internal" / "development" / "Reply_Status_Development_Exercise.pdf"


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_like(doc, text, size=16):
    return add_para(doc, text, size=size, bold=True, space_after=12, space_before=6)


def add_label(doc, label):
    return add_para(doc, label, size=12, bold=True, space_after=4, space_before=10)


def add_body_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.35)
    run = p.add_run(str(text).strip())
    set_run_font(run, size=11)
    return p


def add_checkbox_line(doc, text):
    return add_para(doc, f"☐  {text}", size=12, space_after=4, space_before=2)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "12"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "666666"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_row_height(row, cm_height: float):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(cm_height * 567)))  # approx twips
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)


def add_comment_box(doc, height_cm: float = 4.0, *, label: str | None = "Comments"):
    if label:
        add_label(doc, label)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_border(cell, val="single", sz="12", color="888888")
    set_row_height(table.rows[0], height_cm)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(" ")
    set_run_font(run, size=12)
    for _ in range(3):
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(10)
        run2 = p2.add_run(" ")
        set_run_font(run2, size=12)
    add_para(doc, "", size=8, space_after=4)


def clear_core_props(doc: Document) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "Pilot Annotation Exercise: Reply Status"
    props.subject = ""
    props.keywords = ""
    props.category = ""
    props.comments = ""


def build() -> Path:
    df = pd.read_excel(SRC, sheet_name="CODING")
    assert len(df) == 15

    doc = Document()
    clear_core_props(doc)

    section = doc.sections[0]
    # A4
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # -------- Title page --------
    for _ in range(2):
        doc.add_paragraph()

    add_para(
        doc,
        "Pilot Annotation Exercise",
        size=22,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_para(
        doc,
        "Reply Status in Spanish Parliamentary Control Sessions",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )
    add_para(
        doc,
        "A short practice exercise for expert readers",
        size=12,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=24,
    )

    add_para(
        doc,
        (
            "This is a coding-guide development exercise. It is not the final "
            "annotation of the study. There are no right answers. The purpose is to "
            "identify unclear cases, missing rules, and wording that creates uncertainty."
        ),
        size=12,
        space_after=12,
    )
    add_para(
        doc,
        (
            "Please work independently. Disagreement between experts is expected and "
            "useful. After finishing the cases, complete the short feedback page at "
            "the end of this booklet."
        ),
        size=12,
        space_after=12,
    )
    add_para(doc, "Number of cases: 15", size=12, bold=True, space_after=6)
    add_para(
        doc,
        "Estimated time: about 60–90 minutes, depending on case length.",
        size=12,
        space_after=18,
    )

    doc.add_page_break()

    # -------- Instructions --------
    add_heading_like(doc, "Instructions", size=18)

    add_heading_like(doc, "What you are judging", size=14)
    add_para(
        doc,
        (
            "For each case you will see three texts: the registered written question, "
            "the questioner’s opening oral turn, and the Prime Minister’s first response."
        ),
        size=12,
        space_after=8,
    )
    add_para(
        doc,
        (
            "Your task is to judge whether the Prime Minister’s first response answers "
            "the oral question. Use the registered question only as institutional "
            "context when the spoken turn is elliptical or refers back to it. The "
            "object of judgement is the oral question, as answered in the first response."
        ),
        size=12,
        space_after=10,
    )

    add_heading_like(doc, "Categories", size=14)

    add_para(doc, "Explicit reply", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "The response clearly fills the information request in the oral question. "
            "A clear yes or no, a concrete factual answer, or a clear denial of the "
            "asked proposition all count. Rejecting a premise can still be an explicit "
            "reply if the response addresses the ask."
        ),
        size=12,
        space_after=8,
    )

    add_para(doc, "Intermediate reply", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "The response engages the question but leaves the requested content "
            "incomplete, heavily hedged, only partial, or only implied. Answering "
            "only one part of a multi-part oral question usually belongs here."
        ),
        size=12,
        space_after=8,
    )

    add_para(doc, "Non-reply", size=12, bold=True, space_after=2)
    add_para(
        doc,
        (
            "The response does not answer the oral question: it changes the subject, "
            "attacks the questioner without answering, answers a different question, "
            "or offers only procedure or a general government record that never "
            "meets the ask."
        ),
        size=12,
        space_after=10,
    )

    add_heading_like(doc, "Important principles", size=14)
    bullets = [
        "Political agreement with the speaker is irrelevant.",
        "Whether you find the answer persuasive, honest, or good policy is irrelevant.",
        "Criticism of the questioner is allowed; if a clear answer is also present, that can still be an explicit reply.",
        "Disagreement between experts is allowed and informative.",
        "If a case feels close to a boundary, mark it as borderline and use the comments box.",
        "Judge only the first response printed here, not later exchanges.",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(b)
        set_run_font(run, size=12)

    add_para(doc, "", size=12, space_after=6)
    add_para(
        doc,
        (
            "How to fill each case: tick one classification box, tick borderline "
            "yes or no, and write brief comments when the decision was difficult."
        ),
        size=12,
        space_after=8,
    )

    # -------- Cases --------
    for i, row in df.iterrows():
        doc.add_page_break()
        case_no = int(i) + 1
        add_heading_like(doc, f"Case {case_no}", size=18)

        add_label(doc, "Registered question")
        add_body_block(doc, row["registered_question"])

        add_label(doc, "Oral question")
        add_body_block(doc, row["Q1"])

        add_label(doc, "First response")
        add_body_block(doc, row["R1"])

        # Always start the fill-in form on a fresh page so the boxes are never orphaned
        doc.add_page_break()
        add_heading_like(doc, f"Case {case_no} — Your judgement", size=16)
        add_para(
            doc,
            "Tick one classification, say whether the case felt borderline, and add comments if useful.",
            size=11,
            italic=True,
            space_after=10,
        )

        add_label(doc, "Classification")
        add_checkbox_line(doc, "Explicit reply")
        add_checkbox_line(doc, "Intermediate reply")
        add_checkbox_line(doc, "Non-reply")

        add_label(doc, "Borderline")
        add_checkbox_line(doc, "Yes")
        add_checkbox_line(doc, "No")

        add_comment_box(doc, height_cm=7.5)

    # -------- Feedback --------
    doc.add_page_break()
    add_heading_like(doc, "Feedback on the coding guide", size=18)
    add_para(
        doc,
        (
            "Please answer briefly after completing the fifteen cases. This feedback "
            "is for improving the coding guide only."
        ),
        size=12,
        space_after=12,
    )

    questions = [
        "1. Was the category decision clear overall?",
        "2. What rule, if any, felt missing?",
        "3. What wording in the instructions caused uncertainty?",
        "4. Would another expert in your field likely make the same decisions? Why or why not?",
        "5. Optional: which case numbers were hardest, and why?",
    ]
    for q in questions:
        add_para(doc, q, size=12, bold=True, space_after=6, space_before=10)
        add_comment_box(doc, height_cm=3.2, label=None)

    add_para(doc, "", size=12, space_after=8)
    add_para(
        doc,
        "Thank you for your time and judgement.",
        size=12,
        italic=True,
        space_after=6,
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    path = build()
    print(path)
