"""Cuaderno de discusión de los cinco desacuerdos del ejercicio inicial.

No revisa la guía de criterios. No produce una categoría de consenso.
Documento visible en español académico de España.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "_internal" / "data_private" / "development" / "out_of_pool" / "DEVELOPMENT_PACKET_BLINDED.csv"
OUT_DOCX = ROOT / "_internal" / "development" / "REVISION_DE_DESACUERDOS_DESARROLLO.docx"

DISAGREEMENT_CASES = (1, 6, 7, 12, 13)
INCLUDE_REGISTERED = {1, 12, 13}

UNIT_BY_CASE = {
    1: "dev_94c479d43b2df7ec",
    6: "dev_85e6a5b707885fe4",
    7: "dev_705bd6dfd843b995",
    12: "dev_36b002c772ec6e62",
    13: "dev_aa60c52418bc3e00",
}

DANIEL = {
    1: "ausencia de respuesta",
    6: "respuesta parcial o intermedia",
    7: "respuesta parcial o intermedia",
    12: "respuesta parcial o intermedia",
    13: "respuesta explícita",
}
JOSE = {
    1: "respuesta parcial o intermedia",
    6: "ausencia de respuesta",
    7: "ausencia de respuesta",
    12: "ausencia de respuesta",
    13: "ausencia de respuesta",
}


def set_run_font(run, *, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p



def add_heading_like(doc, text, size=16, space_before=10, page_break_before=False):
    return add_para(
        doc,
        text,
        size=size,
        bold=True,
        space_after=10,
        space_before=space_before,
        page_break_before=page_break_before,
    )


def add_label(doc, label, *, size=12):
    return add_para(doc, label, size=size, bold=True, space_after=4, space_before=10)


def add_body_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Cm(0.25)
    run = p.add_run(str(text).strip())
    set_run_font(run, size=11)
    return p


def add_checkbox_line(doc, text):
    return add_para(doc, f"☐  {text}", size=12, space_after=6, space_before=2)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "12"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "666666"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_row_height(row, cm_height: float):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(cm_height * 567)))
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def add_keep_table(doc, paragraphs: list[tuple[str, dict]]) -> None:
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_border(cell, val="nil", sz="0", color="FFFFFF")
    cell.text = ""
    first = True
    for text, opts in paragraphs:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(opts.get("after", 4))
        p.paragraph_format.space_before = Pt(opts.get("before", 0))
        set_run_font(
            p.add_run(text),
            size=opts.get("size", 12),
            bold=opts.get("bold", False),
            italic=opts.get("italic", False),
        )
    add_para(doc, "", size=6, space_after=4)


def add_comment_box(doc, height_cm: float = 3.2, *, label: str | None = None):
    if label:
        add_label(doc, label)
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_border(cell, val="single", sz="12", color="888888")
    set_row_height(table.rows[0], height_cm)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    set_run_font(p.add_run(" "), size=12)
    for _ in range(3):
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(8)
        set_run_font(p2.add_run(" "), size=12)
    add_para(doc, "", size=6, space_after=4)


def clear_core_props(doc: Document) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "Revisión conjunta de casos dudosos"
    props.subject = ""
    props.keywords = ""
    props.category = ""
    props.comments = ""


def add_front_matter(doc: Document) -> None:
    add_para(doc, "", size=12, space_after=18)
    add_para(
        doc,
        "REVISIÓN CONJUNTA DE CASOS DUDOSOS",
        size=20,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    add_para(
        doc,
        "Tipo de respuesta en las sesiones de control parlamentario",
        size=13,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )
    add_para(
        doc,
        (
            "En el primer ejercicio, ambos análisis coincidieron en diez de los "
            "quince casos. El propósito de este documento es examinar los cinco "
            "casos en los que las valoraciones iniciales difirieron. No se trata "
            "de decidir quién tenía razón. Se trata de identificar qué distinción "
            "o qué criterio de decisión conviene expresar con más claridad antes "
            "de la siguiente fase."
        ),
        size=12,
        space_after=12,
    )
    add_para(
        doc,
        (
            "La mayor parte de los desacuerdos se sitúa en la frontera entre una "
            "respuesta parcial o intermedia y una ausencia de respuesta. Conviene "
            "prestar atención, por tanto, a si basta con hablar del mismo asunto "
            "para que exista una respuesta, o si la intervención debe aportar "
            "información que resuelva al menos una parte de la cuestión planteada. "
            "Esa distinción no está zanjada: es precisamente lo que hay que discutir."
        ),
        size=12,
        space_after=12,
    )
    add_para(
        doc,
        (
            "Las valoraciones iniciales se formularon de manera independiente. "
            "Esta lectura conjunta se abre solo después de haber recibido ambas "
            "devoluciones. El documento sirve para revisar los criterios; no "
            "establece una interpretación definitiva ni un patrón de referencia."
        ),
        size=12,
        space_after=10,
    )


def add_discussion_questions(doc: Document, *, extra_q13: bool = False) -> None:
    add_heading_like(doc, "Preguntas para la discusión", size=14, space_before=8)
    questions = [
        "1. ¿Cuál considera que es la cuestión principal que debe responderse?",
        (
            "2. ¿Qué fragmento concreto de la respuesta considera relevante "
            "para decidir si se responde o no a esa cuestión?"
        ),
        (
            "3. ¿La respuesta aporta información que resuelve total o parcialmente "
            "la cuestión, o se limita a hablar de un asunto relacionado?"
        ),
        (
            "4. ¿Qué criterio debería añadirse o aclararse en la guía para que "
            "este caso resulte más fácil de clasificar?"
        ),
    ]
    if extra_q13:
        questions.append(
            "5. ¿Estabais identificando la misma cuestión principal al realizar "
            "la primera valoración?"
        )
        questions.append(
            "6. Después de discutirlo, ¿consideráis que hace falta modificar "
            "alguna regla?"
        )
    else:
        questions.append(
            "5. Después de discutirlo, ¿consideráis que hace falta modificar "
            "alguna regla?"
        )

    for q in questions[:-1]:
        add_para(doc, q, size=12, bold=True, space_after=4, space_before=6)
        add_comment_box(doc, height_cm=3.0)

    add_keep_table(
        doc,
        [
            (questions[-1], {"after": 8, "size": 12, "bold": True, "before": 6}),
            ("☐  Sí", {"after": 4}),
            ("☐  No", {"after": 10}),
            (
                "Interpretación acordada para orientar la guía (opcional). "
                "Si la registráis, servirá solo para precisar los criterios. "
                "No constituye una interpretación de referencia ni una decisión "
                "definitiva sobre el caso.",
                {"after": 8, "size": 11, "italic": True, "before": 6},
            ),
            ("☐  Respuesta explícita", {"after": 4}),
            ("☐  Respuesta parcial o intermedia", {"after": 4}),
            ("☐  Ausencia de respuesta", {"after": 4}),
        ],
    )


def add_case(doc: Document, case_no: int, row) -> None:
    add_heading_like(
        doc, f"CASO {case_no}", size=18, space_before=0, page_break_before=True
    )

    if case_no in INCLUDE_REGISTERED:
        add_label(doc, "Pregunta registrada")
        add_body_block(doc, row["registered_question"])
        add_para(
            doc,
            (
                "Se incluye la pregunta registrada porque el turno oral no basta "
                "por sí solo para identificar la cuestión, o porque la respuesta "
                "parece referirse a ella."
            ),
            size=10,
            italic=True,
            space_after=8,
        )

    add_label(doc, "Pregunta formulada")
    add_body_block(doc, row["Q1"])

    add_label(doc, "Primera respuesta del Presidente del Gobierno")
    add_body_block(doc, row["R1"])

    add_heading_like(doc, "Valoraciones iniciales", size=14, space_before=8)
    add_para(doc, f"Daniel: {DANIEL[case_no]}.", size=12, space_after=4)
    add_para(doc, f"Jose Jaime: {JOSE[case_no]}.", size=12, space_after=10)

    add_discussion_questions(doc, extra_q13=(case_no == 13))


def add_general_discussion(doc: Document) -> None:
    add_heading_like(
        doc,
        "CRITERIOS QUE DEBEMOS ACLARAR",
        size=18,
        space_before=0,
        page_break_before=True,
    )
    add_para(
        doc,
        (
            "Después de revisar los cinco casos, anoten las precisiones que "
            "consideren necesarias. No hace falta formular aún una regla cerrada; "
            "basta con señalar dónde falla la guía actual."
        ),
        size=12,
        space_after=12,
    )
    prompts = [
        (
            "A. ¿Cuándo una respuesta general sobre el mismo tema constituye "
            "realmente una respuesta?"
        ),
        (
            "B. ¿Cuándo una respuesta relacionada con la pregunta debe "
            "considerarse parcial y cuándo debe considerarse ausencia de respuesta?"
        ),
        (
            "C. Si la respuesta rechaza una premisa de la pregunta, ¿en qué "
            "circunstancias eso constituye una respuesta explícita?"
        ),
        (
            "D. En una intervención con varias preguntas, ¿qué debe ocurrir "
            "para considerar que existe una respuesta explícita?"
        ),
        (
            "E. ¿Hay algún elemento del contexto político que pueda estar "
            "influyendo en la interpretación lingüística? ¿Cómo podemos reducir "
            "ese efecto sin alterar el texto?"
        ),
        "F. ¿Qué ejemplos adicionales debería contener la guía?",
    ]
    for prompt in prompts:
        add_para(doc, prompt, size=12, bold=True, space_after=4, space_before=8)
        add_comment_box(doc, height_cm=3.6)


def build() -> Path:
    pkt = pd.read_csv(SRC)
    by_id = {str(r["unit_id"]): r for _, r in pkt.iterrows()}

    doc = Document()
    clear_core_props(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    add_front_matter(doc)
    for case_no in DISAGREEMENT_CASES:
        uid = UNIT_BY_CASE[case_no]
        add_case(doc, case_no, by_id[uid])
    add_general_discussion(doc)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build())
