"""Ingest independent Calibration Round 1 Word returns and compute agreement.

Does not alter the Round-1 draw, seed, or codebook.
Does not impute missing or multi-marked reply_status choices.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document

from scse.calibration_metrics import ES_TO_EN, summarize_round1
from scse.paths import DATA_PRIVATE, PROJECT_ROOT

ROOT = PROJECT_ROOT
PACKET = DATA_PRIVATE / "calibration" / "round1" / "CALIBRATION_ROUND1_PACKET_BLINDED.csv"
OUT_PRIV = DATA_PRIVATE / "calibration" / "round1"
RET_DIR = ROOT / "_internal" / "calibration_round1" / "returns"
REPORT_MD = ROOT / "_internal" / "reports" / "CALIBRATION_ROUND1_AGREEMENT.md"
REPORT_JSON = ROOT / "_internal" / "reports" / "CALIBRATION_ROUND1_AGREEMENT.json"

LABELS = [
    ("respuesta explícita", "Respuesta explícita"),
    ("respuesta parcial o intermedia", "Respuesta parcial o intermedia"),
    ("ausencia de respuesta", "Ausencia de respuesta"),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def is_marked(line: str) -> bool:
    s = line.strip()
    if re.match(r"^[xX]\s+", s):
        return True
    return s.startswith(("☑", "☒", "✓", "✔"))


def parse_case_block(text: str, case_no: int) -> dict:
    lines = [ln.rstrip() for ln in text.splitlines()]
    reply_marks: dict[str, bool] = {}
    dudoso_marks: dict[str, bool] = {}
    obs_lines: list[str] = []
    mode = None
    for ln in lines:
        low = ln.strip().lower()
        if low.startswith("tu valoración"):
            mode = "reply"
            continue
        if "dudoso" in low:
            mode = "dudoso"
            continue
        if low.startswith("observaciones"):
            mode = "obs"
            continue
        if mode == "reply":
            for canon, needle in LABELS:
                if needle.lower() in low:
                    reply_marks[canon] = is_marked(ln)
        elif mode == "dudoso":
            if re.search(r"^[xX☐☑☒✓✔]\s*sí\b", low) or re.search(
                r"^[xX☐☑☒✓✔]\s*si\b", low
            ):
                dudoso_marks["sí"] = is_marked(ln)
            elif re.search(r"^[xX☐☑☒✓✔]\s*no\b", low):
                dudoso_marks["no"] = is_marked(ln)
        elif mode == "obs":
            if set(ln.strip()) <= {"_", " ", "─", "-", ""}:
                continue
            if ln.strip():
                obs_lines.append(ln.strip())
    selected_reply = [k for k, v in reply_marks.items() if v]
    selected_dudoso = [k for k, v in dudoso_marks.items() if v]
    issues: list[str] = []
    if len(selected_reply) == 0:
        issues.append("missing_reply_status")
    elif len(selected_reply) > 1:
        issues.append("multiple_reply_status")
    if len(selected_dudoso) == 0:
        issues.append("missing_dudoso")
    elif len(selected_dudoso) > 1:
        issues.append("multiple_dudoso")
    return {
        "case_no": case_no,
        "reply_status": selected_reply[0] if len(selected_reply) == 1 else None,
        "reply_status_all_marked": selected_reply,
        "dudoso": selected_dudoso[0] if len(selected_dudoso) == 1 else None,
        "dudoso_all_marked": selected_dudoso,
        "observaciones": "\n".join(obs_lines).strip(),
        "issues": issues,
    }


def extract_doc(path: Path) -> list[dict]:
    doc = Document(path)
    case_nos: list[int] = []
    for p in doc.paragraphs:
        m = re.match(r"^CASO\s+(\d+)\s*$", p.text.strip())
        if m:
            case_nos.append(int(m.group(1)))
    if len(doc.tables) != 20 or case_nos != list(range(1, 21)):
        raise RuntimeError(
            f"bad structure {path.name}: cases={case_nos} tables={len(doc.tables)}"
        )
    return [
        parse_case_block(table.cell(0, 0).text, case_no)
        for case_no, table in zip(case_nos, doc.tables)
    ]


def ingest(daniel_docx: Path, jose_docx: Path) -> dict:
    RET_DIR.mkdir(parents=True, exist_ok=True)
    OUT_PRIV.mkdir(parents=True, exist_ok=True)

    arch = {}
    for key, src in (("daniel", daniel_docx), ("jose_jaime", jose_docx)):
        dest = RET_DIR / f"calibracion_ronda1_{key}_returned.docx"
        shutil.copy2(src, dest)
        arch[key] = {
            "path": str(dest.relative_to(ROOT)),
            "sha256": sha256(dest),
            "source_filename": src.name,
        }

    extracted = {
        "daniel": extract_doc(daniel_docx),
        "jose_jaime": extract_doc(jose_docx),
    }
    pkt = pd.read_csv(PACKET).sort_values("case_no")
    unit_by_case = dict(zip(pkt.case_no.astype(int), pkt.unit_id.astype(str)))

    rows = []
    for i in range(20):
        d = extracted["daniel"][i]
        j = extracted["jose_jaime"][i]
        rows.append(
            {
                "case_no": i + 1,
                "unit_id": unit_by_case[i + 1],
                "daniel_reply_status": d["reply_status"],
                "jose_jaime_reply_status": j["reply_status"],
                "agreement_reply_status": bool(
                    d["reply_status"]
                    and j["reply_status"]
                    and d["reply_status"] == j["reply_status"]
                ),
                "pair_complete": d["reply_status"] is not None
                and j["reply_status"] is not None,
                "daniel_dudoso": d["dudoso"],
                "jose_jaime_dudoso": j["dudoso"],
                "daniel_observaciones": d["observaciones"],
                "jose_jaime_observaciones": j["observaciones"],
                "daniel_issues": ";".join(d["issues"]),
                "jose_jaime_issues": ";".join(j["issues"]),
                "daniel_all_marked": "|".join(d["reply_status_all_marked"]),
                "jose_jaime_all_marked": "|".join(j["reply_status_all_marked"]),
            }
        )
    df = pd.DataFrame(rows)
    complete = df[df.pair_complete]
    metrics = summarize_round1(
        complete["daniel_reply_status"].tolist(),
        complete["jose_jaime_reply_status"].tolist(),
    )
    metrics["n_complete_pairs"] = int(len(complete))
    metrics["n_incomplete_pairs"] = int((~df.pair_complete).sum())
    metrics["n_agree_complete"] = int(complete["agreement_reply_status"].sum())
    metrics["n_disagree_complete"] = int((~complete["agreement_reply_status"]).sum())
    metrics["disagree_case_nos_complete"] = (
        complete.loc[~complete["agreement_reply_status"], "case_no"]
        .astype(int)
        .tolist()
    )
    metrics["incomplete_case_nos"] = (
        df.loc[~df.pair_complete, "case_no"].astype(int).tolist()
    )

    csv_path = OUT_PRIV / "CALIBRATION_ROUND1_CODES.csv"
    df.to_csv(csv_path, index=False)

    incomplete_detail = []
    for _, r in df.loc[~df.pair_complete].iterrows():
        incomplete_detail.append(
            {
                "case_no": int(r.case_no),
                "daniel_reply_status": r.daniel_reply_status,
                "jose_jaime_reply_status": r.jose_jaime_reply_status
                if pd.notna(r.jose_jaime_reply_status)
                else None,
                "daniel_issues": r.daniel_issues,
                "jose_jaime_issues": r.jose_jaime_issues,
                "jose_jaime_all_marked": r.jose_jaime_all_marked,
            }
        )

    payload = {
        "phase": "5A",
        "task": "calibration_round1_returns_ingested",
        "recorded_on": date.today().isoformat(),
        "codebook_version": "0.3.0",
        "seed": 2071684612,
        "blinded_packet_csv_sha256": sha256(PACKET),
        "returned_files": arch,
        "n_cases": 20,
        "n_complete_pairs": metrics["n_complete_pairs"],
        "n_incomplete_pairs": metrics["n_incomplete_pairs"],
        "incomplete_case_nos": metrics["incomplete_case_nos"],
        "incomplete_detail": incomplete_detail,
        "agreement_complete_pairs": {
            "n_agree": metrics["n_agree_complete"],
            "n_disagree": metrics["n_disagree_complete"],
            "raw_agreement": metrics["raw_agreement"],
            "cohen_kappa_3way": metrics["cohen_kappa_3way"],
            "krippendorff_alpha_nominal": metrics["krippendorff_alpha_nominal"],
            "binary_explicit_raw_agreement": metrics[
                "binary_explicit_raw_agreement"
            ],
            "cohen_kappa_explicit_vs_rest": metrics[
                "cohen_kappa_explicit_vs_rest"
            ],
            "confusion_matrix_3way": metrics["confusion_matrix_3way"],
            "disagree_case_nos": metrics["disagree_case_nos_complete"],
        },
        "cases": rows,
        "note": (
            "Exploratory calibration diagnostics only; not final validation; "
            "not gold; codebook v0.3 not frozen for main study."
        ),
    }
    (OUT_PRIV / "CALIBRATION_ROUND1_RETURNS.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    slim = {k: v for k, v in payload.items() if k != "cases"}
    slim["cases_slim"] = [
        {
            "case_no": r["case_no"],
            "unit_id": r["unit_id"],
            "daniel_reply_status": r["daniel_reply_status"],
            "jose_jaime_reply_status": r["jose_jaime_reply_status"],
            "agreement_reply_status": r["agreement_reply_status"],
            "pair_complete": r["pair_complete"],
            "daniel_dudoso": r["daniel_dudoso"],
            "jose_jaime_dudoso": r["jose_jaime_dudoso"],
            "daniel_has_observaciones": bool(r["daniel_observaciones"]),
            "jose_jaime_has_observaciones": bool(r["jose_jaime_observaciones"]),
            "daniel_issues": r["daniel_issues"],
            "jose_jaime_issues": r["jose_jaime_issues"],
        }
        for r in rows
    ]
    REPORT_JSON.write_text(
        json.dumps(slim, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return {
        "metrics": metrics,
        "csv": str(csv_path.relative_to(ROOT)),
        "report_md": str(REPORT_MD.relative_to(ROOT)),
        "en_labels_used_in_matrix": list(ES_TO_EN.values()),
        **arch,
    }


def main() -> int:
    import argparse

    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--daniel", type=Path, required=True)
    p.add_argument("--jose-jaime", type=Path, required=True)
    args = p.parse_args()
    out = ingest(args.daniel, args.jose_jaime)
    print(json.dumps(out["metrics"], indent=2, ensure_ascii=False))
    print("wrote", out["csv"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
