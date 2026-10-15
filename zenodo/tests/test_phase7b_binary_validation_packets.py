"""Phase 7B: issued binary validation materials; sample not redrawn."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd
from docx import Document

from scse.paths import DATA_PRIVATE, PROJECT_ROOT, PROTOCOL_DIR

VAL = DATA_PRIVATE / "validation" / "binary_v1"
HUMAN = PROJECT_ROOT / "_internal" / "binary_validation"
INSTRUMENT = PROTOCOL_DIR / "direct_resolution_codebook_v1.0-validation.md"


def test_validation_instrument_frozen_hash():
    assert INSTRUMENT.is_file()
    expected = json.loads(
        (VAL / "BINARY_VALIDATION_INSTRUMENT.json").read_text(encoding="utf-8")
    )["instrument_sha256"]
    got = hashlib.sha256(INSTRUMENT.read_bytes()).hexdigest()
    assert got == expected
    assert expected.startswith("bd4350472908d95b")


def test_sample_not_redrawn():
    man = json.loads((VAL / "BINARY_VALIDATION_MANIFEST.json").read_text())
    pkt = pd.read_csv(VAL / "BINARY_VALIDATION_N12_PACKET_BLINDED.csv")
    assert list(pkt["unit_id"]) == man["unit_ids_selected"]
    assert len(pkt) == 12
    assert man.get("not_calibration_round_3") is True
    assert man["annotator_docx_generated"] is True


def test_human_packets_exist_and_binary_only():
    orient = HUMAN / "01_ORIENTACION_CONJUNTA_RESOLUCION_DIRECTA.docx"
    daniel = HUMAN / "02_validacion_binaria_Daniel.docx"
    jose = HUMAN / "02_validacion_binaria_Jose_Jaime.docx"
    email = HUMAN / "EMAIL_VALIDACION_BINARIA_FINAL.txt"
    assert orient.is_file() and daniel.is_file() and jose.is_file() and email.is_file()
    td = "\n".join(p.text for p in Document(daniel).paragraphs)
    tj = "\n".join(p.text for p in Document(jose).paragraphs)
    assert "Sí — resolución directa" in td and "No — no resolución directa" in td
    assert "Respuesta explícita" not in td and "Respuesta parcial" not in td
    assert "Respuesta explícita" not in tj
    assert "cal1_" not in td and "kappa" not in td.lower()
    assert td.replace("Para Daniel", "Para X") == tj.replace("Para Jose Jaime", "Para X")


def test_source_texts_intact():
    pkt = pd.read_csv(VAL / "BINARY_VALIDATION_N12_PACKET_BLINDED.csv")
    daniel = Document(HUMAN / "02_validacion_binaria_Daniel.docx")
    full = "\n".join(p.text for p in daniel.paragraphs)
    for _, row in pkt.iterrows():
        assert str(row["registered_question"]).strip() in full
        assert str(row["Q1"]).strip() in full
        assert str(row["R1"]).strip() in full


def test_gate_locked_and_no_main_xiv():
    gate = json.loads((VAL / "BINARY_VALIDATION_PROSPECTIVE_GATE.json").read_text())
    assert gate["kappa_threshold"] == 0.67
    assert gate["raw_agreement_threshold_n"] == 10
    assert gate["automatic_further_round_on_fail"] is False
    pool = pd.read_parquet(DATA_PRIVATE / "eligible_pool_full.parquet")
    assert "reply_status" not in pool.columns
    assert "direct_resolution" not in pool.columns
    assert (PROJECT_ROOT / "_internal" / "reports" / "BINARY_VALIDATION_PACKET_AUDIT.md").is_file()


def test_email_sequence_clear():
    email = (HUMAN / "EMAIL_VALIDACION_BINARIA_FINAL.txt").read_text(encoding="utf-8")
    assert "orientación" in email.lower() or "orientacion" in email.lower()
    assert "12" in email
    assert "kappa" not in email.lower()
    assert "Gate C" not in email
